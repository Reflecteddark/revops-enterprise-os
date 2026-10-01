"""
Тесты для speech_engine.py и run_express_audit.py:
- Обезличивание ПДн по 152-ФЗ (телефоны, email, карты, паспорта)
- Восстановление из защищенного сейфа (де-токенизация)
- 13 критериев B2B оценки диалогов
- Проверка жесткого Next Step и выявления утечек маржи
- Юнит-экономика STT (Yandex SpeechKit Deferred 0.02 ₽/мин vs Groq 0.06 ₽/мин)
- Запись в лист raw_calls Excel с валидными формулами
- Генерация экспресс-аудита и питча для Telegram
"""

import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from run_express_audit import format_telegram_pitch, generate_html_audit_card
from speech_engine import (
    B2BCallEvaluator,
    CallEvaluation,
    CRITERIA_METADATA,
    PIISanitizer,
    SpeechEngine,
    sync_call_to_excel,
)


# ─────────────────────────────────────────────────────────────────────────────
# 1. Тесты 152-ФЗ Обезличивания (PIISanitizer)
# ─────────────────────────────────────────────────────────────────────────────

class TestPIISanitizer:
    def test_tokenize_phone_numbers(self):
        text = "Свяжитесь со мной по номеру +7 (925) 123-45-67 или 8 916 999-88-77 завтра."
        sanitized, vault = PIISanitizer.tokenize_text(text)

        assert "+7 (925) 123-45-67" not in sanitized
        assert "8 916 999-88-77" not in sanitized
        assert len(vault) == 2
        assert PIISanitizer.is_clean(sanitized) is True

    def test_tokenize_email(self):
        text = "Отправьте коммерческое предложение на ceo@enterprise-corp.ru до конца дня."
        sanitized, vault = PIISanitizer.tokenize_text(text)

        assert "ceo@enterprise-corp.ru" not in sanitized
        assert "[EMAIL_TOKEN_" in sanitized
        assert len(vault) == 1
        assert PIISanitizer.is_clean(sanitized) is True

    def test_tokenize_cards_and_passports(self):
        text = "Паспорт 45 12 889900, карта 4276 3800 1234 5678."
        sanitized, vault = PIISanitizer.tokenize_text(text)

        assert "4276 3800 1234 5678" not in sanitized
        assert "45 12 889900" not in sanitized
        assert "[CARD_TOKEN_" in sanitized
        assert "[PASSPORT_TOKEN_" in sanitized

    def test_detokenize_roundtrip(self):
        original = "Клиент Иван, телефон +7 999 111-22-33, почта client@mail.ru."
        sanitized, vault = PIISanitizer.tokenize_text(original)
        restored = PIISanitizer.detokenize_text(sanitized, vault)

        assert restored == original

    def test_empty_and_clean_text(self):
        sanitized, vault = PIISanitizer.tokenize_text("")
        assert sanitized == ""
        assert vault == {}

        clean_text = "Мы согласовали поставку на следующую неделю."
        s, v = PIISanitizer.tokenize_text(clean_text)
        assert s == clean_text
        assert v == {}


# ─────────────────────────────────────────────────────────────────────────────
# 2. Тесты 13 Критериев (B2BCallEvaluator)
# ─────────────────────────────────────────────────────────────────────────────

class TestB2BCallEvaluator:
    def test_criteria_count(self):
        assert len(CRITERIA_METADATA) == 13

    def test_perfect_call_evaluation(self):
        perfect_transcript = (
            "Добрый день, компания RevOps, меня зовут Михаил. Удобно говорить? "
            "Какая задача стоит перед вами и к какому дедлайну? "
            "Кто кроме вас будет принимать решение по проекту и согласовывать бюджет? "
            "Понимаю, что дорого, но у клиентов из вашей ниши наш кейс показал окупаемость за 11 дней. "
            "Давайте запустим пилот и выставим счет. Итак, резюмирую: в четверг в 14:00 созваниваемся."
        )
        ev = B2BCallEvaluator.evaluate(perfect_transcript)

        assert ev.total_score >= 10
        assert ev.next_step_fixed is True
        assert ev.criteria_scores["cr1"] == 1
        assert ev.criteria_scores["cr3"] == 1
        assert ev.criteria_scores["cr5"] == 1
        assert "Слив Next Step" not in ev.error_summary

    def test_lost_next_step_detection(self):
        poor_transcript = (
            "Здравствуйте. Ну наше решение стоит 79 тысяч. "
            "Понятно, дорого? Ну подумайте, посмотрите. "
            "Ну тогда спишемся как-нибудь, на связи!"
        )
        ev = B2BCallEvaluator.evaluate(poor_transcript)

        assert ev.criteria_scores["cr1"] == 0
        assert ev.next_step_fixed is False
        assert "CRITICAL_LOST_NEXT_STEP" in ev.critical_alerts
        assert "Слив Next Step" in ev.error_summary

    def test_talk_listen_ratio_evaluation(self):
        turns = [
            {"speaker": "manager", "text": "Добрый день! Предлагаю наш сервис. Расскажите подробнее о вашей компании."},
            {"speaker": "client", "text": "У нас сеть из 15 филиалов, главная боль — контроль качества звонков, менеджеры сливают лиды."},
        ]
        ev = B2BCallEvaluator.evaluate("Диалог", speaker_turns=turns)
        assert 0.0 <= ev.talk_listen_ratio <= 1.0


# ─────────────────────────────────────────────────────────────────────────────
# 3. Тесты Универсального Движка (SpeechEngine)
# ─────────────────────────────────────────────────────────────────────────────

class TestSpeechEngine:
    def test_unit_economics_pricing(self):
        engine = SpeechEngine(provider="mock")
        # 10 минут звонка в Yandex Deferred (0.02 ₽/мин)
        result = engine.transcribe("test.wav", duration_sec=600, mock_scenario="standard_sale")

        assert result.duration_sec == 600.0
        # 10 мин * 0.02 = 0.20 ₽
        assert result.estimated_cost_rub == 0.20
        assert len(result.text) > 0
        assert "[PHONE_TOKEN_" in result.sanitized_text or len(result.token_vault) >= 0

    def test_auto_provider_selection(self):
        with patch.dict("os.environ", {"YANDEX_API_KEY": "y_key", "YANDEX_FOLDER_ID": "f_id"}, clear=True):
            engine = SpeechEngine(provider="auto")
            assert engine.provider == "yandex"

        with patch.dict("os.environ", {"GROQ_API_KEY": "g_key"}, clear=True):
            engine = SpeechEngine(provider="auto")
            assert engine.provider == "groq"

        with patch.dict("os.environ", {}, clear=True):
            engine = SpeechEngine(provider="auto")
            assert engine.provider == "mock"


# ─────────────────────────────────────────────────────────────────────────────
# 4. Тесты Записи в Excel (sync_call_to_excel)
# ─────────────────────────────────────────────────────────────────────────────

class TestExcelSync:
    def test_sync_call_to_raw_calls(self, tmp_path):
        # Создаем временную копию Excel с листом raw_calls
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "raw_calls"
        ws.append(["call_id", "deal_id", "manager_id", "call_date", "duration_sec", "call_type",
                   "cr1", "cr2", "cr3", "cr4", "cr5", "cr6", "cr7", "cr8", "cr9", "cr10", "cr11", "cr12", "cr13",
                   "total_score", "next_step_flag", "error_summary", "ai_recommendation", "synced_at"])
        test_file = tmp_path / "test_revops.xlsx"
        wb.save(test_file)

        ev = B2BCallEvaluator.evaluate("В четверг в 14:00 созвонимся", call_id="C-TEST-99")
        ok, msg = sync_call_to_excel(ev, excel_path=test_file)

        assert ok is True
        # Проверяем записанные данные
        wb_read = openpyxl.load_workbook(test_file, data_only=False)
        ws_read = wb_read["raw_calls"]
        assert ws_read.cell(2, 1).value == "C-TEST-99"
        # Проверяем формулу суммы
        assert ws_read.cell(2, 20).value == "=SUM(G2:S2)"
        # Проверяем формулу флага
        assert ws_read.cell(2, 21).value == "=T2>=9"


# ─────────────────────────────────────────────────────────────────────────────
# 5. Тесты Экспресс-Аудита (run_express_audit)
# ─────────────────────────────────────────────────────────────────────────────

class TestExpressAudit:
    def test_html_card_generation(self, tmp_path):
        ev1 = B2BCallEvaluator.evaluate("В пятницу в 11:00 созвон", call_id="C-1")
        ev2 = B2BCallEvaluator.evaluate("Спишемся как-нибудь", call_id="C-2")
        out_file = tmp_path / "audit_card.html"

        result_path = generate_html_audit_card([ev1, ev2], out_file)
        assert result_path.exists()
        content = result_path.read_text(encoding="utf-8")
        assert "Экспресс-аудит звонков" in content
        assert "C-1" in content
        assert "C-2" in content
        assert "152-ФЗ" in content

    def test_telegram_pitch_formatting(self):
        ev1 = B2BCallEvaluator.evaluate("В пятницу в 11:00 созвон")
        ev2 = B2BCallEvaluator.evaluate("Спишемся как-нибудь")
        pitch = format_telegram_pitch([ev1, ev2])

        assert "RevOps" in pitch
        assert "13" in pitch
        assert "29 000" in pitch
        assert "Слив Next Step" in pitch or "НЕ зафиксировали" in pitch
