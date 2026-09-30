import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

target_files = [
    'RevOps Platform V17.6 (RBAC Production Suite).xlsx',
    'C:/Users/strel/.gemini/antigravity/brain/6a2f6569-0174-4758-bbb9-311ec84e9bf3/RevOps_Enterprise_OS_V17.6_Final_Remediated.xlsx'
]

f_a1 = '=IF(ISNUMBER(SEARCH("CEO", \'⚙️ Настройки\'!$B$10)), "📋 СТРАТЕГИЧЕСКИЙ ПУЛЬТ СОБСТВЕННИКА: «5 РЕШЕНИЙ МЕСЯЦА»", IF(ISNUMBER(SEARCH("KAM", \'⚙️ Настройки\'!$B$10)), "📋 ЛИЧНЫЙ ACTION-LIST МЕНЕДЖЕРА НА СЕГОДНЯ (ПЕТРОВ А.)", IF(ISNUMBER(SEARCH("SDR", \'⚙️ Настройки\'!$B$10)), "📞 ПУЛЬТ ТЕЛЕМАРКЕТОЛОГА: «СКРИПТЫ И ГОРЯЧИЕ ЗВОНКИ»", IF(ISNUMBER(SEARCH("CMO", \'⚙️ Настройки\'!$B$10)), "📋 ПУЛЬТ ОПТИМИЗАЦИИ ТРАФИКА: «ТОЧКИ СЛИВА БЮДЖЕТА»", IF(ISNUMBER(SEARCH("CFO", \'⚙️ Настройки\'!$B$10)), "📋 ПЛАТЕЖНЫЙ КАЛЕНДАРЬ И ДЕБИТОРКА: «СРОЧНЫЕ СЧЕТА»", IF(OR(ISNUMBER(SEARCH("Admin", \'⚙️ Настройки\'!$B$10)), ISNUMBER(SEARCH("Администр", \'⚙️ Настройки\'!$B$10))), "🔧 ПУЛЬТ АДМИНИСТРАТОРА: «МОНИТОРИНГ И НАСТРОЙКИ СИСТЕМЫ»", "📋 ОПЕРАТИВНЫЙ ПУЛЬТ РОПа: «15 МИНУТ В ДЕНЬ»"))))))'
f_a7 = '="🎯 " & IF(ISNUMBER(SEARCH("CEO", \'⚙️ Настройки\'!$B$10)), "ТОП-5 СТРАТЕГИЧЕСКИХ РЕШЕНИЙ ДЛЯ ЗАЩИТЫ КАПИТАЛА БИЗНЕСА", IF(ISNUMBER(SEARCH("KAM", \'⚙️ Настройки\'!$B$10)), "МОЙ ПЕРСОНАЛЬНЫЙ СПИСОК ГОРЯЩИХ СДЕЛОК НА СЕГОДНЯ", IF(ISNUMBER(SEARCH("SDR", \'⚙️ Настройки\'!$B$10)), "ТОП-5 ЗВОНКОВ В РИСКЕ И СРОЧНАЯ КВАЛИФИКАЦИЯ ЛИДОВ", IF(ISNUMBER(SEARCH("CMO", \'⚙️ Настройки\'!$B$10)), "ТОП-5 ДЕЙСТВИЙ ПО ОПТИМИЗАЦИИ РЕКЛАМНЫХ КАМПАНИЙ", IF(ISNUMBER(SEARCH("CFO", \'⚙️ Настройки\'!$B$10)), "ТОП-5 СРОЧНЫХ ДЕЙСТВИЙ ПО СБОРУ ДЕБИТОРКИ И ЗАЩИТЕ КАССЫ", IF(OR(ISNUMBER(SEARCH("Admin", \'⚙️ Настройки\'!$B$10)), ISNUMBER(SEARCH("Администр", \'⚙️ Настройки\'!$B$10))), "ТОП-5 СИСТЕМНЫХ АЛЕРТОВ И ПРОВЕРКА ЦЕЛОСТНОСТИ REVOPS", "ТОП-5 СРОЧНЫХ ДЕЙСТВИЙ ДЛЯ РОПа НА СЕГОДНЯ (DAILY ACTION LIST)"))))))'

f_b9 = '=HYPERLINK("#\'Action_Center\'!A1", "Сделка D-104 (ООО «Вектор Плюс»)")'
f_b10 = '=HYPERLINK("#\'🎙️ ИИ_Аудит\'!A1", "Звонок C-505 / Сделка D-109 (ООО «ПромКомплект»)")'
f_b11 = '=HYPERLINK("#\'🎙️ ИИ_Аудит\'!A1", "Звонок C-503 / Сделка D-106 (ООО «СнабСервис»)")'
f_b12 = '=HYPERLINK("#\'Action_Center\'!A1", "Сделка D-107 (ООО «ТехноПром»)")'
f_b13 = '=HYPERLINK("#\'Action_Center\'!A1", "Сделка D-103 (ООО «ИнвестХолдинг»)")'

for fpath in target_files:
    print(f'Syncing fixes to {fpath}...')
    wb = openpyxl.load_workbook(fpath)
    
    # 1. 📋 Пульт_РОПа_15_Минут
    if '📋 Пульт_РОПа_15_Минут' in wb.sheetnames:
        ws = wb['📋 Пульт_РОПа_15_Минут']
        ws['A1'] = f_a1
        ws['A7'] = f_a7
        ws['B9'] = f_b9
        ws['B10'] = f_b10
        ws['B11'] = f_b11
        ws['B12'] = f_b12
        ws['B13'] = f_b13

    # 2. Versions
    if '🧪 QA_Suite' in wb.sheetnames:
        wb['🧪 QA_Suite']['A1'] = "118 РЕГРЕССИОННЫХ ТЕСТОВ (REVOPS ENTERPRISE V17.6)"
    if '⚙️ Настройки' in wb.sheetnames:
        wb['⚙️ Настройки']['C9'] = "Enterprise Multi-Tenant V17.6"
    if '🌐 Мультиканальная_Атрибуция' in wb.sheetnames:
        wb['🌐 Мультиканальная_Атрибуция']['A1'] = "🌐 МУЛЬТИКАНАЛЬНАЯ АТРИБУЦИЯ ВЫРУЧКИ: W-SHAPED REVENUE ENGINE (V17.6)"
    if '💳 Финансы_и_AI_Дожим' in wb.sheetnames:
        wb['💳 Финансы_и_AI_Дожим']['A1'] = "💳 ФИНАНСОВЫЙ КОНТУР, REVERSE ETL И АГЕНТНЫЙ AI-ДОЖИМ (V17.6 BLOCK I)"
    if '👥 Мотивация_ОП' in wb.sheetnames:
        wb['👥 Мотивация_ОП']['A1'] = "👥 СИСТЕМА ДИНАМИЧЕСКОЙ МОТИВАЦИИ И PAYROLL ОП (REVOPS V17.6 BLOCK II)"
    if '🔮 Симулятор_Роста' in wb.sheetnames:
        wb['🔮 Симулятор_Роста']['A1'] = "🔮 СЦЕНАРНЫЙ СИМУЛЯТОР ВЫРУЧКИ И МАСШТАБИРОВАНИЯ (WHAT-IF REVENUE SANDBOX V17.6)"
    if '📊 Юнит_Экономика' in wb.sheetnames:
        wb['📊 Юнит_Экономика']['A1'] = "📊 ЮНИТ-ЭКОНОМИКА, КОГОРТНЫЙ NRR И LTV/CAC (V17.6)"

    # 3. ⚡ Пульс_Компании
    if '⚡ Пульс_Компании' in wb.sheetnames:
        ws = wb['⚡ Пульс_Компании']
        ws['F9'] = 'План выручки'
        ws['G9'] = "='⚙️ Настройки'!$B$9"
        ws['F10'] = 'Факт Closed-Won'
        ws['G10'] = "=A4"
        ws['F11'] = 'Run-Rate прогноз'
        ws['G11'] = "=C4"
        ws['F12'] = 'Взвешенный прогноз'
        ws['G12'] = "=D4"
        ws['F13'] = 'Активный пайплайн'
        ws['G13'] = "=E4"

    # 4. 📄 Executive_OnePager formatting
    if '📄 Executive_OnePager' in wb.sheetnames:
        ws = wb['📄 Executive_OnePager']
        ws['A5'].number_format = '#,##0 "₽"'
        ws['B5'].number_format = '0.0%'
        ws['C5'].number_format = '#,##0 "₽"'
        ws['D5'].number_format = '#,##0 "₽"'
        ws['E5'].number_format = '#,##0 "₽"'
        ws['F5'].number_format = '#,##0'
        ws['G5'].number_format = '0.0%'
        ws['H5'].number_format = '0.0%'

    wb.save(fpath)
    print(f'Saved {fpath} successfully!')

print('All local XLSX files updated with fixes!')
