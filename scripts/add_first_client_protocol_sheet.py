"""
Скрипт добавления листа '🚀 Протокол_Первого_Клиента' в RevOps_Founder_Workspace_Planner.xlsx
Создает исчерпывающий пошаговый регламент действий фаундера: от входящей заявки
до бесплатного аудита 3 звонков, дожима на платный пилот (45к-150к) и запуска абонентки.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")


def add_first_client_protocol(file_path: Path):
    if not file_path.exists():
        print(f"File not found: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet_name = "🚀 Протокол_Первого_Клиента"

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True

    # Цветовая палитра
    c_navy_dark = "0F172A"       # #0F172A
    c_navy_header = "1E293B"     # #1E293B
    c_blue_primary = "1E3A8A"    # #1E3A8A
    c_blue_light = "EFF6FF"      # #EFF6FF
    c_emerald = "059669"         # #059669
    c_emerald_light = "ECFDF5"   # #ECFDF5
    c_amber = "D97706"           # #D97706
    c_amber_light = "FFFBEB"     # #FFFBEB
    c_rose = "E11D48"            # #E11D48
    c_rose_light = "FFF1F2"      # #FFF1F2
    c_light_bg = "F8FAFC"        # #F8FAFC
    c_border = "CBD5E1"          # #CBD5E1

    fill_navy = PatternFill(start_color=c_navy_dark, end_color=c_navy_dark, fill_type="solid")
    fill_header = PatternFill(start_color=c_blue_primary, end_color=c_blue_primary, fill_type="solid")
    fill_sub_hdr = PatternFill(start_color=c_navy_header, end_color=c_navy_header, fill_type="solid")
    fill_zebra = PatternFill(start_color=c_light_bg, end_color=c_light_bg, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_blue_chip = PatternFill(start_color=c_blue_light, end_color=c_blue_light, fill_type="solid")
    fill_emerald_chip = PatternFill(start_color=c_emerald_light, end_color=c_emerald_light, fill_type="solid")
    fill_amber_chip = PatternFill(start_color=c_amber_light, end_color=c_amber_light, fill_type="solid")

    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Segoe UI", size=9.5, italic=True, color="94A3B8")
    font_header = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    font_sec = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
    font_reg = Font(name="Segoe UI", size=9, color="334155")
    font_chip_blue = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
    font_chip_emerald = Font(name="Segoe UI", size=9, bold=True, color="059669")
    font_chip_amber = Font(name="Segoe UI", size=9, bold=True, color="D97706")
    font_code = Font(name="Consolas", size=8.5, color="0F172A")

    thin_border = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border)
    )

    # 1. Заголовок
    ws.merge_cells("A1:H1")
    ws["A1"] = "🚀 RevOps Enterprise OS | Пошаговый Регламент Ведения Первого Клиента (End-to-End SOP)"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    ws["A2"] = "Полная инструкция действий основателя: от входящей заявки до аудита 3 звонков за 24ч, закрытия пилота 45к-150к ₽ и абонентского сопровождения"
    ws["A2"].font = font_sub
    ws["A2"].fill = fill_navy
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # 2. Ключевые метрики и KPI (Строки 4-5)
    ws.merge_cells("A4:B4")
    ws["A4"] = "⚡ СКОРОСТЬ ПЕРВОГО ОТВЕТА"
    ws["A4"].font = font_bold
    ws["A4"].fill = fill_blue_chip
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("C4:D4")
    ws["C4"] = "⏱️ ФОРМАТ ПЕРВОГО ДЕМО"
    ws["C4"].font = font_bold
    ws["C4"].fill = fill_blue_chip
    ws["C4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("E4:F4")
    ws["E4"] = "🎯 ГЛАВНЫЙ ВХОДНОЙ ОФФЕР"
    ws["E4"].font = font_bold
    ws["E4"].fill = fill_emerald_chip
    ws["E4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("G4:H4")
    ws["G4"] = "💰 ЦЕЛЕВОЙ ЧЕК ПИЛОТА"
    ws["G4"].font = font_bold
    ws["G4"].fill = fill_amber_chip
    ws["G4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A5:B5")
    ws["A5"] = "< 10 минут (в WhatsApp/TG)"
    ws["A5"].font = font_chip_blue
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C5:D5")
    ws["C5"] = "Ровно 20 минут в Zoom"
    ws["C5"].font = font_chip_blue
    ws["C5"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("E5:F5")
    ws["E5"] = "Тест 3 звонков за 24ч (0 ₽)"
    ws["E5"].font = font_chip_emerald
    ws["E5"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("G5:H5")
    ws["G5"] = "45 000 – 150 000 ₽"
    ws["G5"].font = font_chip_amber
    ws["G5"].alignment = Alignment(horizontal="center", vertical="center")

    for r in range(4, 6):
        ws.row_dimensions[r].height = 24
        for c in range(1, 9):
            ws.cell(r, c).border = thin_border

    # 3. Основная таблица: 12 Шагов
    ws.merge_cells("A7:H7")
    ws["A7"] = "📋 ПОШАГОВЫЙ БОЕВОЙ АЛГОРИТМ: 12 ШАГОВ ВЕДЕНИЯ КЛИЕНТА"
    ws["A7"].font = font_sec
    ws["A7"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[7].height = 26

    headers = [
        "Шаг", "Фаза воронки", "Срок / Дедлайн", "Что конкретно делает фаундер",
        "Инструмент / Где лежит", "Что сказать / Написать (скрипт)", "Критерий успеха (DoD)", "Опасная ошибка (Чего нельзя делать)"
    ]

    ws.row_dimensions[8].height = 28
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(8, col_idx, h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    steps_data = [
        (
            "Шаг 1", "Входящий лид", "0 – 10 минут",
            "Лид упал с сайта ai-rop.ru на почту info@ai-rop.ru или в Telegram @dm1918. Моментально пишем клиенту. Проводим экспресс-квалификацию по 3 вопросам.",
            "Почта / Telegram @dm1918",
            "«[Имя], добрый день! Получил вашу заявку на ai-rop.ru. Подскажите, сколько у вас менеджеров в отделе, какой средний чек и какую CRM используете?»",
            "Получен ответ: подтверждено от 3 менеджеров и средний чек от 50 000 ₽.",
            "Ждать больше 15 минут; отправлять громоздкие КП и коммерческие предложения в первом сообщении."
        ),
        (
            "Шаг 2", "Назначение демо", "В течение 1 часа",
            "Предложить 2 конкретных слота времени на выбор для 20-минутного Zoom-разбора. Скинуть ссылку на встречу и внести в календарь.",
            "Google Календарь / Zoom / Яндекс.Телемост",
            "«Отлично! Предлагаю созвониться на 20 минут в Zoom: покажу на пальцах, как ИИ слушает 100% звонков и находит потери. Завтра удобно в 11:30 или в 15:00?»",
            "Встреча подтверждена клиентом и стоит в календаре.",
            "Спрашивать: «Когда вам удобно?» — это размазывает договоренность."
        ),
        (
            "Шаг 3", "20-мин Презентация", "День 2 (строго 20 мин)",
            "Провести созвон строго по готовому Word-сценарию. Открыть сайт ai-rop.ru, Excel-калькулятор (3 числа) и Google Таблицу (13 стандартов B2B).",
            "Сценарий_20_минутной_Презентации_RevOps.docx",
            "Вести разговор по файлу: 3 мин боль -> 4 мин Калькулятор потерь -> 6 мин Живой аудит -> 3 мин 152-ФЗ -> 4 мин Оффер на тест 3 звонков.",
            "Клиент в шоке от суммы потенциальных потерь и согласен на тест 3 звонков.",
            "Пытаться продать годовой софт за 300к на первой встрече. Цель первой встречи — ТОЛЬКО получить 3 звонка!"
        ),
        (
            "Шаг 4", "Забор 3 звонков", "Сразу после созвона",
            "Отправить персональную ссылку на облачную папку для загрузки аудио (Яндекс.Диск/Google Drive) или предложить прислать файлы прямо в Telegram.",
            "Яндекс.Диск / Telegram @dm1918",
            "«[Имя], спасибо за созвон! Вот ссылка на закрытую папку: [Ссылка]. Загрузите 3 любых звонка за вчера (1 успешный и 2 слитых). Ровно через 24ч верну аудит».",
            "3 аудиофайла (.mp3 или .wav) сохранены на вашем компьютере.",
            "Давать сложные технические инструкции по API телефонии. Брать аудиофайлы простым архивом!"
        ),
        (
            "Шаг 5", "Защита 152-ФЗ", "По запросу клиента",
            "Если клиент или его СБ сомневаются в конфиденциальности — мгновенно отправить подписанный NDA и памятку по 152-ФЗ.",
            "Папка «152 ФЗ» на Рабочем столе",
            "«Мы работаем строго по 152-ФЗ РФ. Перед началом подписываем двусторонний NDA. Данные деперсонализируются в защищенном контуре. Направляю проект NDA».",
            "NDA подписан с обеих сторон (скан или ЭДО).",
            "Спорить со службой безопасности; говорить, что данные отправляются на зарубежные сервера."
        ),
        (
            "Шаг 6", "Fulfillment (Запуск ИИ)", "В течение 4–6 часов",
            "Запустить Start_Sales_Calls_AI.bat (Faster-Whisper на 8000, DeepSeek-R1 на 11434, n8n на 5678). Прогнать 3 аудиозаписи через транскрибацию и скоринг.",
            "Start_Sales_Calls_AI.bat, transcriber_service",
            "Автоматический процесс: получение расшифровки с тайм-кодами и скоринга по 13 критериям B2B.",
            "Получены точные транскрипты и баллы соответствия по всем 13 стандартам.",
            "Оставлять сырой нечитаемый текст без разбивки на реплики менеджера и клиента."
        ),
        (
            "Шаг 7", "Сборка клиентского отчета", "В течение 8 часов",
            "Сделать копию мастер-таблицы Google Sheets. Заполнить вкладку «ИИ_Аудит»: вставить цитаты, выставить зеленые/красные чипы, рассчитать недополученную маржу.",
            "Google Sheets Master (gid 852624872)",
            "Форматирование в корпоративном стиле: чипы ошибок, выделение желтым зон роста, формулы ROI.",
            "Готовая брендированная ссылка на Google Таблицу с правами «Просмотр» для клиента.",
            "Использовать неаккуратные шрифты или оставлять ошибки '###' в ячейках."
        ),
        (
            "Шаг 8", "Human-in-the-Loop", "За 2 часа до сдачи",
            "Фаундер лично переслушивает 2 ключевых спорных фрагмента. Формулирует 2 убийственных инсайта (например: 'Менеджер дал скидку 100к без торга').",
            "Google Таблица + блокнот инсайтов",
            "«Инсайт 1: Менеджер Александр не выявил ЛПР и отпустил клиента без даты повторного звонка. Инсайт 2: Слив скидки 15% снизил маржинальность сделки на 240 000 ₽».",
            "Готовы 2 железобетонных факта потерь в деньгах, с которыми собственник не сможет поспорить.",
            "Полагаться на 100% только на авто-генерацию ИИ без личной проверки спикером."
        ),
        (
            "Шаг 9", "Презентация аудита", "Ровно через 24 часа",
            "Созвон на 15 минут. Демонстрация экрана: «Вот ваши 3 звонка, вот где менеджеры слили 400 000 ₽». Показ готовой Google Таблицы.",
            "Google Таблица клиента + Zoom",
            "«[Имя], посмотрите строку 12: клиент спросил про рассрочку, менеджер ответил 'нет' и положил трубку. Здесь компания потеряла 350к. И это только 1 звонок из 3!»",
            "Клиент признает: аудит объективен, проблема реальна и стоит больших денег.",
            "Отправлять ссылку молча в мессенджер без личной демонстрации на созвоне."
        ),
        (
            "Шаг 10", "Закрытие на пилот", "Финальные 5 мин созвона",
            "Презентовать 7-дневный спринт: 100% звонков отдела за 45 000 – 75 000 ₽ (или спринт на месяц за 150 000 ₽) с гарантией окупаемости.",
            "Спецификация пилота / Памятка 7 дней",
            "«Мы отслушали 3 звонка и нашли слив на 400к. За 7 дней мы подключим 100% звонков отдела, дадим дашборд РОПу и спасем от 1 млн ₽. Стоимость пилота — 55 000 ₽. Стартуем?»",
            "Устное «ДА» от собственника и запрос счета/договора.",
            "Делать скидки при первом сомнении. Держать цену твердо, опираясь на ROI калькулятора."
        ),
        (
            "Шаг 11", "Оплата и договор", "В течение 24 часов",
            "Выставить счет от ИП/Самозанятого с QR-кодом. Направить типовой договор оказания услуг. Проконтролировать приход денег.",
            "Банк-клиент (Точка / Тинькофф / Сбер)",
            "«[Имя], направил договор и счет с QR-кодом на почту и в WhatsApp. Как только бухгалтерия проведет оплату, мы за 2 часа запускаем коннектор звонков».",
            "Деньги поступили на расчетный счет. Подписан скан договора.",
            "Начинать полноценную работу без предоплаты."
        ),
        (
            "Шаг 12", "Сопровождение и LTV", "Дни 1 – 7 пилота",
            "Ежедневная выгрузка звонков, генерация утренней сводки в Telegram для РОПа в 09:00. В пятницу — итоговый отчет и продажа абонентской подписки (60к-120к/мес).",
            "Telegram-бот @RevOps_Super_Audit_Bot",
            "«За 7 дней мы разобрали 420 звонков, подняли показатель жесткого следующего шага с 32% до 78%, спасли 3 сделки на 1.2 млн ₽. Переходим на помесячный контракт?»",
            "Подписан долгосрочный контракт на регулярную аналитику (MRR от 60 000 ₽).",
            "Оставлять клиента без ежедневного контакта; не показывать динамику улучшений."
        )
    ]

    current_row = 9
    for row_idx, data in enumerate(steps_data, start=1):
        ws.row_dimensions[current_row].height = 48
        bg = fill_zebra if row_idx % 2 == 1 else fill_white
        
        for col_idx, text in enumerate(data, start=1):
            cell = ws.cell(current_row, col_idx, text)
            cell.font = font_reg
            cell.fill = bg
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            
            if col_idx == 1:
                cell.font = font_bold
                cell.fill = fill_blue_chip
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.font = font_bold
            elif col_idx == 3:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                cell.font = font_code
            elif col_idx == 8:
                cell.font = Font(name="Segoe UI", size=8.5, color=c_rose, bold=True)
                cell.fill = PatternFill(start_color=c_rose_light, end_color=c_rose_light, fill_type="solid")

        current_row += 1

    # 4. Блок шаблонов быстрых сообщений
    current_row += 1
    ws.merge_cells(f"A{current_row}:H{current_row}")
    ws.cell(current_row, 1, "💬 ШАБЛОНЫ БЫСТРЫХ СООБЩЕНИЙ ДЛЯ КОПИРОВАНИЯ (FAST COPY-PASTE)")
    ws.cell(current_row, 1).font = font_sec
    ws.cell(current_row, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[current_row].height = 26
    current_row += 1

    templates_hdr = ["Код", "Событие", "Канал", "Готовый текст сообщения для отправки клиенту", "", "", "", ""]
    ws.row_dimensions[current_row].height = 24
    for c_i, h_t in enumerate(["Код", "Событие", "Канал", "Готовый текст сообщения для клиента"], start=1):
        cell = ws.cell(current_row, c_i, h_t)
        cell.font = font_header
        cell.fill = fill_sub_hdr
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"D{current_row}:H{current_row}")
    current_row += 1

    templates_data = [
        (
            "MSG-01", "Первый ответ на заявку с сайта", "Telegram / WhatsApp",
            "«[Имя], приветствую! Получил вашу заявку на сервисе ai-rop.ru по аудиту звонков отдела продаж. Подскажите, пожалуйста: 1) Сколько менеджеров сейчас в линии? 2) Какой средний чек сделки? И какую CRM используете (amoCRM или Битрикс24)? Хочу подготовить релевантные примеры под вашу нишу перед созвоном»."
        ),
        (
            "MSG-02", "Напоминание за 1 час до Zoom-демо", "Telegram / WhatsApp",
            "«[Имя], добрый день! Напоминаю, что через 1 час (в [Время]) у нас запланирован 20-минутный созвон в Zoom. Ссылка на встречу: [Ссылка]. Покажу на живых цифрах, как работает ИИ-аудит и сколько выручки теряется в звонках. До встречи!»"
        ),
        (
            "MSG-03", "Инструкция по передаче 3 звонков", "Telegram / WhatsApp",
            "«[Имя], спасибо за встречу! Как договорились, жду от вас 3 любых аудиозаписи за вчера (один успешный звонок и два с сорвавшимися сделками). Загрузить можно сюда: [Ссылка на Яндекс.Диск] либо скинуть аудиофайлами прямо в этот чат. Завтра в [Время] верну готовую аналитику по 13 критериям»."
        ),
        (
            "MSG-04", "Уведомление о готовности аудита", "Telegram / WhatsApp",
            "«[Имя], добрый день! Аудит ваших 3 звонков полностью готов. Нашли 2 критические системные ошибки менеджеров, из-за которых сорвались сделки на суммарно ~[Сумма] ₽. Собрал всё в персональную Google Таблицу. Предлагаю созвониться на 15 минут — выведу экран и покажу конкретные цитаты и точки слива. Удобно в [Время]?»"
        ),
        (
            "MSG-05", "Отправка счета на 7-дневный пилот", "Email + Telegram",
            "«[Имя], направил вам договор и счет на 7-дневный пилот RevOps (сумма 55 000 ₽). В стоимость входит анализ 100% звонков вашего отдела, ежедневная утренняя сводка РОПу в Telegram и оцифровка точек потерь. Оплатить можно по реквизитам или по QR-коду в счете. Стартуем сразу после подтверждения!»"
        )
    ]

    for code, event, channel, text in templates_data:
        ws.row_dimensions[current_row].height = 42
        c1 = ws.cell(current_row, 1, code)
        c1.font = font_bold
        c1.fill = fill_blue_chip
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.border = thin_border

        c2 = ws.cell(current_row, 2, event)
        c2.font = font_bold
        c2.alignment = Alignment(vertical="center")
        c2.border = thin_border

        c3 = ws.cell(current_row, 3, channel)
        c3.font = font_reg
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c3.border = thin_border

        ws.merge_cells(f"D{current_row}:H{current_row}")
        c4 = ws.cell(current_row, 4, text)
        c4.font = font_code
        c4.alignment = Alignment(vertical="center", wrap_text=True)
        for c_idx in range(4, 9):
            ws.cell(current_row, c_idx).border = thin_border

        current_row += 1

    # Ширины колонок
    col_widths = {
        "A": 9,    # Шаг / Код
        "B": 18,   # Фаза воронки
        "C": 16,   # Срок
        "D": 38,   # Что делает
        "E": 28,   # Где лежит
        "F": 46,   # Что сказать
        "G": 30,   # DoD
        "H": 28    # Ошибка
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Сохранение
    wb.save(file_path)
    print(f"Sheet successfully added to: {file_path}")


if __name__ == "__main__":
    targets = [
        Path(r"C:\Users\strel\Desktop\RevOps_Founder_Workspace_Planner.xlsx"),
        Path(r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\docs\RevOps_Founder_Workspace_Planner.xlsx"),
        Path(r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\presentation\RevOps_Founder_Workspace_Planner.xlsx")
    ]
    for target in targets:
        add_first_client_protocol(target)
