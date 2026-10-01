"""
RevOps Enterprise OS V17.6 — Synchronizer & Manager for Founder's Workspace Planner.
Модуль для автоматического синхронного ведения рабочей книги фаундера
(RevOps_Founder_Workspace_Planner.xlsx).
"""

from __future__ import annotations

import datetime
import os
import shutil
import sys
from pathlib import Path
from typing import Optional

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PLANNER_PATH = Path("presentation/RevOps_Founder_Workspace_Planner.xlsx")
DOCS_PATH = Path("docs/RevOps_Founder_Workspace_Planner.xlsx")
DESKTOP_PATH = Path(os.path.expanduser(r"~\Desktop\RevOps_Founder_Workspace_Planner.xlsx"))

c_border = "CBD5E1"
thin_border = Border(
    left=Side(style="thin", color=c_border),
    right=Side(style="thin", color=c_border),
    top=Side(style="thin", color=c_border),
    bottom=Side(style="thin", color=c_border)
)

STATUS_FILLS = {
    "✅ ВЫПОЛНЕНО": PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid"),
    "⏳ В РАБОТЕ": PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid"),
    "📋 БЭКЛОГ": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
}

PRIORITY_FONTS = {
    "P0": Font(name="Segoe UI", size=10, bold=True, color="DC2626"),
    "P1": Font(name="Segoe UI", size=10, bold=True, color="D97706"),
    "P2": Font(name="Segoe UI", size=10, color="64748B"),
}


def _save_and_replicate(wb: openpyxl.Workbook):
    """Сохраняет файл в presentation, docs и на Рабочий стол."""
    PLANNER_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PLANNER_PATH)

    DOCS_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(PLANNER_PATH, DOCS_PATH)

    if DESKTOP_PATH.parent.exists():
        shutil.copyfile(PLANNER_PATH, DESKTOP_PATH)


def update_task_status(task_id: str, new_status: str, notes: Optional[str] = None) -> bool:
    """Обновляет статус задачи по её ID (например, T-08 -> ✅ ВЫПОЛНЕНО)."""
    wb = openpyxl.load_workbook(PLANNER_PATH)
    ws = wb["🎯 Спринты_и_Задачи"]

    found = False
    for r in range(4, ws.max_row + 1):
        cell_id = ws.cell(r, 1).value
        if cell_id and str(cell_id).strip().upper() == task_id.strip().upper():
            # Колонка 6: Статус
            ws.cell(r, 6, new_status)
            ws.cell(r, 6).font = Font(name="Segoe UI", size=10, bold=True)
            if new_status in STATUS_FILLS:
                ws.cell(r, 6).fill = STATUS_FILLS[new_status]
            
            # Колонка 8: Заметки
            if notes:
                ws.cell(r, 8, notes)
            found = True
            break

    if found:
        _save_and_replicate(wb)
        print(f"Статус задачи {task_id} обновлен на: {new_status}")
    else:
        print(f"Задача {task_id} не найдена в таск-трекере.")
    return found


def add_task(
    stream: str,
    title: str,
    priority: str = "P1",
    deadline: str = "",
    status: str = "⏳ В РАБОТЕ",
    artifact: str = "",
    notes: str = "",
) -> str:
    """Добавляет новую задачу в таск-трекер фаундера."""
    wb = openpyxl.load_workbook(PLANNER_PATH)
    ws = wb["🎯 Спринты_и_Задачи"]

    # Находим следующую свободную строку и определяем ID
    next_row = 4
    max_id = 0
    while ws.cell(next_row, 1).value is not None:
        raw_id = str(ws.cell(next_row, 1).value)
        if raw_id.startswith("T-"):
            try:
                num = int(raw_id.split("-")[1])
                max_id = max(max_id, num)
            except ValueError:
                pass
        next_row += 1

    new_task_id = f"T-{max_id + 1:02d}"
    if not deadline:
        deadline = datetime.datetime.now().strftime("%d.%m.%Y")

    ws.row_dimensions[next_row].height = 24
    ws.cell(next_row, 1, new_task_id).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(next_row, 2, stream).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(next_row, 3, title)
    
    p_cell = ws.cell(next_row, 4, priority)
    p_cell.alignment = Alignment(horizontal="center", vertical="center")
    p_cell.font = PRIORITY_FONTS.get(priority, Font(name="Segoe UI", size=10))

    ws.cell(next_row, 5, deadline).alignment = Alignment(horizontal="center", vertical="center")
    
    st_cell = ws.cell(next_row, 6, status)
    st_cell.alignment = Alignment(horizontal="center", vertical="center")
    st_cell.font = Font(name="Segoe UI", size=10, bold=True)
    if status in STATUS_FILLS:
        st_cell.fill = STATUS_FILLS[status]

    ws.cell(next_row, 7, artifact)
    ws.cell(next_row, 8, notes)

    for col in range(1, 9):
        ws.cell(next_row, col).border = thin_border
        if col != 4 and col != 6:
            ws.cell(next_row, col).font = Font(name="Segoe UI", size=10, color="334155")

    _save_and_replicate(wb)
    print(f"Добавлена новая задача {new_task_id}: {title} [{status}]")
    return new_task_id


def add_partner(
    company: str,
    stack: str = "amoCRM",
    city: str = "Москва",
    contact: str = "",
    status: str = "В плане на контакт",
    revshare: str = "25%",
    potential: str = "10 активных",
    next_step: str = "Отправить оффер",
) -> None:
    """Добавляет нового интегратора в пайплайн партнеров."""
    wb = openpyxl.load_workbook(PLANNER_PATH)
    ws = wb["🤝 Пайплайн_Интеграторов"]

    # Ищем первую пустую строку (где компания не указана)
    target_row = None
    for r in range(4, 30):
        if ws.cell(r, 2).value is None or str(ws.cell(r, 2).value).strip() == "":
            target_row = r
            break

    if target_row is None:
        target_row = ws.max_row + 1

    idx = target_row - 3
    ws.row_dimensions[target_row].height = 24
    ws.cell(target_row, 1, idx).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(target_row, 2, company)
    ws.cell(target_row, 3, stack).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(target_row, 4, city)
    ws.cell(target_row, 5, contact)
    
    st_cell = ws.cell(target_row, 6, status)
    st_cell.alignment = Alignment(horizontal="center", vertical="center")
    st_cell.font = Font(name="Segoe UI", size=10, bold=True)
    st_cell.fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")

    ws.cell(target_row, 7, revshare)
    ws.cell(target_row, 8, potential)
    ws.cell(target_row, 9, next_step)

    for c in range(1, 10):
        ws.cell(target_row, c).border = thin_border
        if c != 6:
            ws.cell(target_row, c).font = Font(name="Segoe UI", size=10, color="334155")

    _save_and_replicate(wb)
    print(f"Партнер '{company}' добавлен в строку {target_row}")


def add_client(
    company: str,
    niche: str,
    managers: int,
    check: int,
    crm: str,
    status: str = "Лид на аудит",
    notes: str = "",
) -> None:
    """Добавляет нового клиента в воронку пилотов."""
    wb = openpyxl.load_workbook(PLANNER_PATH)
    ws = wb["🎙️ Клиентские_Пилоты"]

    target_row = None
    for r in range(4, 30):
        if ws.cell(r, 2).value is None or str(ws.cell(r, 2).value).strip() == "":
            target_row = r
            break

    if target_row is None:
        target_row = ws.max_row + 1

    idx = target_row - 3
    ws.row_dimensions[target_row].height = 24
    ws.cell(target_row, 1, idx).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(target_row, 2, company)
    ws.cell(target_row, 3, niche)
    ws.cell(target_row, 4, managers).alignment = Alignment(horizontal="center", vertical="center")
    
    chk_cell = ws.cell(target_row, 5, check)
    chk_cell.number_format = "#,##0 ₽"
    chk_cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.cell(target_row, 6, crm).alignment = Alignment(horizontal="center", vertical="center")
    
    st_cell = ws.cell(target_row, 7, status)
    st_cell.alignment = Alignment(horizontal="center", vertical="center")
    st_cell.font = Font(name="Segoe UI", size=10, bold=True)

    tier = "Старт (49к)" if managers <= 5 else ("Бизнес (79к)" if managers <= 10 else "Enterprise (120к)")
    ws.cell(target_row, 8, "29 000 ₽")
    ws.cell(target_row, 9, tier)
    ws.cell(target_row, 10, notes)

    for c in range(1, 11):
        ws.cell(target_row, c).border = thin_border
        if c != 7:
            ws.cell(target_row, c).font = Font(name="Segoe UI", size=10, color="334155")

    _save_and_replicate(wb)
    print(f"Клиент '{company}' добавлен в воронку пилотов (строка {target_row})")


if __name__ == "__main__":
    print("Проверка модуля sync_founder_planner.py:")
    # Тестовая проверка обновления статуса
    update_task_status("T-06", "✅ ВЫПОЛНЕНО", "Интерактивный B2B лендинг готов, протестирован, сохранен в Git")
