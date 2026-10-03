import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil

sys.stdout.reconfigure(encoding="utf-8")

COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate Navy #0F172A
COLOR_ACCENT = RGBColor(37, 99, 235)      # Royal Blue #2563EB
COLOR_SECONDARY = RGBColor(71, 85, 105)   # Muted Slate #475569
COLOR_SUCCESS = RGBColor(16, 185, 129)    # Emerald Green #10B981
COLOR_MUTED = RGBColor(100, 116, 139)     # Gray #64748B

HEX_BG_LIGHT = "F8FAFC"
HEX_BG_HEADER = "0F172A"
HEX_BORDER = "E2E8F0"
HEX_ACCENT_BG = "EFF6FF"
HEX_SUCCESS_BG = "ECFDF5"

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    header = section.header
    p_head = header.paragraphs[0]
    p_head.text = "REVOPS ENTERPRISE OS V18.0  |  ОФИЦИАЛЬНАЯ ПРЕЗЕНТАЦИЯ ПРОДУКТА"
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.runs[0].font.size = Pt(8)
    p_head.runs[0].font.color.rgb = COLOR_MUTED
    p_head.runs[0].font.name = "Calibri"
    
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.text = "RevOps Architecture  •  100% Речевой ИИ-Аудит  •  Единый Пульт РОПа  •  revops-enterprise.ru"
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.runs[0].font.size = Pt(8)
    p_foot.runs[0].font.color.rgb = COLOR_MUTED
    p_foot.runs[0].font.name = "Calibri"

def add_slide_separator(doc, slide_title, slide_number):
    if slide_number > 1:
        doc.add_page_break()
    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(8)
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run(f"СЛАЙД {slide_number} ИЗ 11")
    r_num.font.name = "Calibri"
    r_num.font.size = Pt(9)
    r_num.font.bold = True
    r_num.font.color.rgb = COLOR_ACCENT
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(slide_title)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY

def build_presentation_docx():
    doc = docx.Document()
    add_header_footer(doc)
    
    # СЛАЙД 1: ОБЛОЖКА
    add_slide_separator(doc, "Добро пожаловать в RevOps Enterprise OS!", 1)
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(12)
    r1 = p1.add_run("Интеллектуальная платформа сквозного контроля выручки, 100% речевого ИИ-аудита диалогов и оперативного управления отделом продаж для B2B-компаний.\n\n"
                    "Готовое промышленное решение для Собственников (CEO) и Руководителей отделов продаж (РОП).")
    r1.font.name = "Calibri"
    r1.font.size = Pt(12)
    r1.font.color.rgb = COLOR_SECONDARY
    
    # СЛАЙД 2: АРХИТЕКТУРА ЭКОСИСТЕМЫ
    add_slide_separator(doc, "Архитектура экосистемы RevOps Core", 2)
    tbl2 = doc.add_table(rows=1, cols=4)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    cols_data = [
        ("AI Speech Inspector", "Собственный сервер речевой аналитики. Автоматически транскрибирует 100% звонков менеджеров и оценивает их по 13 жестким стандартам продаж."),
        ("amoCRM / B24 Reverse ETL", "Двусторонний бесшовный коннектор. За 15 секунд возвращает в карточку сделки текстовый аудит, оценку, теги и задачи Next Step."),
        ("Пульт РОПа & SLA", "Рабочий стол руководителя: 15-минутная утренняя планерка, контроль зависших КП (>48ч), платежный календарь дебиторки (DSO)."),
        ("Telegram War Room Bot", "Мгновенный радар критических инцидентов. Оповещает РОПа за 15 минут о сливе крупной сделки, формирует утренние и вечерние сводки.")
    ]
    for idx, (title, desc) in enumerate(cols_data):
        cell = tbl2.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"📦 {title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_ACCENT
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY
    
    # СЛАЙД 3: ЧЕМ ЗАНИМАЕТСЯ REVOPS OS
    add_slide_separator(doc, "Чем занимается RevOps Enterprise OS?", 3)
    tbl3 = doc.add_table(rows=1, cols=3)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    tri_data = [
        ("🛡️ Оцифровывает и спасает упущенную прибыль", "Находит 4 главные системные утечки отдела продаж (слив на звонках, зависание КП >48ч, брошенные отказники, задержка оплат) и возвращает до 28% упущенной выручки компании."),
        ("🎙️ Слушает и объективно оценивает 100% звонков", "Полностью освобождает РОПа от рутинного прослушивания сотен аудиозаписей. Нейросеть беспристрастно аудирует каждый диалог и дает персональные Coaching Tips менеджерам."),
        ("📊 Дает руководству управляемость и прогноз кассы", "Заменяет субъективные отчеты на единый пульт с прогнозом закрытия месяца (Run-Rate), контролем регламентов SLA и прозрачной формулой динамической мотивации (Payroll).")
    ]
    for idx, (title, desc) in enumerate(tri_data):
        cell = tbl3.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"{title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(10)
        r_d.font.color.rgb = COLOR_SECONDARY
        
    # СЛАЙД 4: КТО НАШИ КЛИЕНТЫ
    add_slide_separator(doc, "Кто наши клиенты и какие задачи мы решаем?", 4)
    tbl4 = doc.add_table(rows=1, cols=2)
    tbl4.alignment = WD_TABLE_ALIGNMENT.CENTER
    ceo_tasks = (
        "👑 СОБСТВЕННИКИ И ГЕНЕРАЛЬНЫЕ ДИРЕКТОРА (CEO)\n\n"
        "1. Прекращаем слив рекламного бюджета на неквалифицированных звонках менеджеров.\n"
        "2. Обеспечиваем математически точный прогноз кассы к концу месяца (Run-Rate).\n"
        "3. Оцифровываем скрытые финансовые потери коммерческого блока в рублях.\n"
        "4. Исключаем зависимость бизнеса от «незаменимых» и токсичных менеджеров-звезд.\n"
        "5. Внедряем честный расчет зарплат: выплата бонусов только за соблюдение регламентов.\n"
        "6. Защищаем компанию от кассовых разрывов через контроль дебиторской задолженности (DSO)."
    )
    rop_tasks = (
        "📋 РУКОВОДИТЕЛИ ОТДЕЛОВ ПРОДАЖ (РОП / КОММЕРЧЕСКИЙ ДИРЕКТОР)\n\n"
        "1. Сокращаем утреннюю планерку с 1.5 часов до 15 минут строго по фактам и цифрам.\n"
        "2. Экономим до 3 часов в день за счет автоматического 100% аудита звонков нейросетью.\n"
        "3. Мгновенно узнаем о срыве крупных сделок и успеваем перехватить клиента за 15 минут.\n"
        "4. Гарантируем 100% контроль железобетонного Next Step (дата и время контакта).\n"
        "5. Получаем готовые Coaching Tips — персональные фразы-подсказки для каждого менеджера.\n"
        "6. Устраняем рутину ручного контроля заполнения карточек и задач в amoCRM."
    )
    for idx, text in enumerate([ceo_tasks, rop_tasks]):
        cell = tbl4.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_PRIMARY
        
    # СЛАЙД 5: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 1 - РЕЧЕВОЙ АУДИТ)
    add_slide_separator(doc, "Как это работает: Модуль 1. Речевой ИИ-Аудит Звонков", 5)
    tbl5 = doc.add_table(rows=1, cols=3)
    tbl5.alignment = WD_TABLE_ALIGNMENT.CENTER
    s5_data = [
        ("Шаг 1. Бесшовный захват аудио из CRM", "Интеграция перехватывает аудиозапись из любой телефонии (Манго, Телфин, Sipuni, Мои Звонки). Звонки 100% менеджеров захватываются автоматически. 0 ручных действий."),
        ("Шаг 2. Скоринг по 13 критериям продаж", "ИИ транскрибирует речь и сопоставляет диалог с жестким регламентом компании. Проверка приветствия, выявления болей, бюджета, дедлайнов и железобетонного Next Step."),
        ("Шаг 3. Формирование Coaching Tip", "Система выставляет оценку 0–100, определяет зону риска (<70) и генерирует конкретную рекомендацию: какую фразу нужно было сказать вместо ошибки.")
    ]
    for idx, (title, desc) in enumerate(s5_data):
        cell = tbl5.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"⚡ {title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_ACCENT
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 6: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 2 - AMOCRM)
    add_slide_separator(doc, "Как это работает: Модуль 2. Двусторонняя Автоматизация amoCRM", 6)
    tbl6 = doc.add_table(rows=1, cols=3)
    tbl6.alignment = WD_TABLE_ALIGNMENT.CENTER
    s6_data = [
        ("Шаг 1. Развернутое саммари в сделку", "Через 15 секунд после звонка в ленте карточки появляется структурированное примечание: пересказ сути договоренностей, баллы по 13 критериям и подсказка РОПа. Заполняется поле «Оценка ИИ»."),
        ("Шаг 2. Автопостановка задач Next Step", "Если менеджер забыл назначить дату перезвона, ИИ сам ставит жесткую задачу до конца дня. Если клиент попросил КП к четвергу в 14:00, робот создаст задачу с точным дедлайном."),
        ("Шаг 3. Умное тегирование и Follow-Up", "Автоматическая маркировка: #Брак_Речи, #Слив_Клиента, #Эталонный_Звонок, #Возражение_Дорого. Генерация готового текста сообщения в WhatsApp/Telegram для отправки клиенту в 1 клик.")
    ]
    for idx, (title, desc) in enumerate(s6_data):
        cell = tbl6.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"⚡ {title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_ACCENT
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 7: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 3 - УПРАВЛЕНИЕ И БЕЗОПАСНОСТЬ)
    add_slide_separator(doc, "Как это работает: Модуль 3. Операционный Пульт и Безопасность", 7)
    tbl7 = doc.add_table(rows=1, cols=3)
    tbl7.alignment = WD_TABLE_ALIGNMENT.CENTER
    s7_data = [
        ("1. Утренний Пульт РОПа (15 минут)", "Оценка темпа месяца (Pacing): идем ли по графику плана выручки. Светофор воронки: выявление зависших сделок на этапе КП (>48ч). ТОП-5 критических рисков выручки для немедленного дожима."),
        ("2. Telegram War Room & Алерты", "Красная кнопка РОПу: при срыве сделки >300 000 ₽ — сигнал в Telegram за 15 минут. Утренняя сводка в 09:00 (рейтинг сейлзов). Вечерний отчет собственнику в 19:00 (касса и прогноз)."),
        ("3. Безопасность и 152-ФЗ РФ", "Полное соблюдение Федерального закона 152-ФЗ: сервера в РФ. Изоляция данных Multi-Tenant по Tenant UUID. Защита интеллектуальной собственности и закрытый контур промптов.")
    ]
    for idx, (title, desc) in enumerate(s7_data):
        cell = tbl7.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"⚡ {title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_ACCENT
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 8: КОМПАНИЯ В ЦИФРАХ
    add_slide_separator(doc, "Платформа RevOps OS в цифрах", 8)
    tbl8 = doc.add_table(rows=2, cols=4)
    tbl8.alignment = WD_TABLE_ALIGNMENT.CENTER
    nums_data = [
        ("100%", "звонков проверяются ИИ без участия человека"),
        ("15 минут", "длительность утренней планерки РОПа по фактам"),
        ("+18–34%", "средний прирост чистой выручки за первые 60 дней"),
        ("< 48 часов", "жесткий норматив нахождения сделки на этапе КП"),
        ("13 стандартов", "объективной оценки диалогов в речевой матрице"),
        ("118 тестов", "автоматической проверки точности QA Suite"),
        ("4 утечки", "финансовых потерь ликвидируются под ключ"),
        ("15 минут", "техническое время подключения нового клиента")
    ]
    for idx, (num, desc) in enumerate(nums_data):
        r_idx = idx // 4
        c_idx = idx % 4
        cell = tbl8.cell(r_idx, c_idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_n = p.add_run(f"{num}\n")
        r_n.bold = True
        r_n.font.name = "Calibri"
        r_n.font.size = Pt(16)
        r_n.font.color.rgb = COLOR_ACCENT
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(8.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 9: ФАКТЫ О REVOPS OS
    add_slide_separator(doc, "Ключевые факты о RevOps Enterprise OS", 9)
    tbl9 = doc.add_table(rows=2, cols=4)
    tbl9.alignment = WD_TABLE_ALIGNMENT.CENTER
    facts_data = [
        ("Единый контур SSOT", "Объединение речевого аудита, CRM-автоматизации и пульта РОПа в один экран."),
        ("Двусторонний обмен CRM", "Прямая интеграция с amoCRM/Битрикс24 без виджетов в браузер сотрудников."),
        ("100% 152-ФЗ РФ Compliant", "Серверная база размещена в аккредитованных дата-центрах РФ."),
        ("Закрытый AI Black Box", "Промпты и скоринговые модели защищены и постоянно дообучаются."),
        ("Быстрый старт за 1 день", "Запуск пилота без остановки работы отдела продаж. Данные уже завтра."),
        ("Адаптация под B2B-нишу", "Калибровка 13 стандартов под специфику опта, услуг или производства."),
        ("Динамический Payroll", "Авторасчет премий с дисконтом 30% за регулярный брак переговоров."),
        ("Полная окупаемость пилота", "Возврат инвестиций в 3–5 раз за счет спасения 1–2 зависших сделок.")
    ]
    for idx, (title, desc) in enumerate(facts_data):
        r_idx = idx // 4
        c_idx = idx % 4
        cell = tbl9.cell(r_idx, c_idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"✅ {title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(8.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 10: ПОЧЕМУ МЫ ДОСТИГЛИ РЕЗУЛЬТАТОВ
    add_slide_separator(doc, "Почему мы достигаем таких результатов? Наши принципы", 10)
    tbl10 = doc.add_table(rows=1, cols=3)
    tbl10.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_data = [
        ("1. Цифры и факты вместо субъективных отчетов", "Менеджеры больше не могут оправдываться «клиент думает». Руководство видит объективную стенограмму разговора, соблюдение этапов и точные таймкоды договоренностей. Решения принимаются на базе реальных данных."),
        ("2. Фокус на живой выручке, а не звонках ради галочки", "Мы оцениваем не бессмысленную длительность звонков, а продвижение сделки к деньгам: квалификацию ЛПР, выявление бюджета, сокращение времени КП и соблюдение сроков оплаты счетов (DSO)."),
        ("3. Обучение команды без стресса и микроменеджмента", "ИИ подсказывает ошибки тактично и сразу после звонка через персональные Coaching Tips. Менеджеры растут в квалификации самостоятельно, а РОП становится наставником, а не контролером.")
    ]
    for idx, (title, desc) in enumerate(p_data):
        cell = tbl10.cell(0, idx)
        set_cell_shading(cell, HEX_BG_LIGHT)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        r_t = p.add_run(f"{title}\n\n")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # СЛАЙД 11: ЗАКЛЮЧИТЕЛЬНЫЙ СЛАЙД
    add_slide_separator(doc, "Готовы оцифровать и усилить ваш отдел продаж?", 11)
    p11 = doc.add_paragraph()
    p11.paragraph_format.space_after = Pt(14)
    r11 = p11.add_run(
        "Запустите 7-дневный тест-драйв: мы подключим вашу CRM и покажем полный ИИ-аудит первых 20 звонков вашей команды уже завтра!\n\n"
        "ПОРЯДОК ЗАПУСКА ПИЛОТНОГО ПРОЕКТА:\n"
        "• Шаг 1 (15 минут): Предоставление доступа к amoCRM и запуск штатного коннектора.\n"
        "• Шаг 2 (24 часа): Нейросеть слушает диалоги и выявляет скрытые финансовые утечки.\n"
        "• Шаг 3 (Результат): Презентация карты потерь на Пульте РОПа и запуск регламентов.\n\n"
        "🌐 Платформа: revops-enterprise.ru  •  📱 Telegram: @revops_founder  •  ⏱️ Развертывание: 15 минут"
    )
    r11.font.name = "Calibri"
    r11.font.size = Pt(11)
    r11.font.color.rgb = COLOR_PRIMARY
    
    # Save Word
    output_docx = r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\docs\Презентация_RevOps_Enterprise_OS_11_Слайдов.docx"
    doc.save(output_docx)
    print(f"Word presentation saved to: {output_docx}")
    
    pres_docx = r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\presentation\Презентация_RevOps_Enterprise_OS_11_Слайдов.docx"
    shutil.copyfile(output_docx, pres_docx)
    print(f"Copied to: {pres_docx}")

if __name__ == "__main__":
    build_presentation_docx()
