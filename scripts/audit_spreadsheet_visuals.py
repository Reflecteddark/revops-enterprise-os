import sys
import openpyxl
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

file_path = Path(r"C:\Users\strel\Desktop\RevOps Platform\Презентация\Презентационная_Таблица_RevOps.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=False)

print(f"Auditing workbook: {file_path}")
print(f"Sheets found ({len(wb.sheetnames)}): {wb.sheetnames}\n")

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"--- Sheet: {sheetname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    
    # Check column widths
    cols_with_width = {}
    for col_letter, dim in ws.column_dimensions.items():
        if dim.width:
            cols_with_width[col_letter] = dim.width
    print(f"  Defined column widths: {cols_with_width}")
    
    # Check potential text clipping issues
    potential_clipping = []
    wrap_issues = []
    
    for r in range(1, ws.max_row + 1):
        row_dim = ws.row_dimensions.get(r)
        r_height = row_dim.height if row_dim else None
        
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            if val is None:
                continue
            val_str = str(val)
            
            # Check if cell is part of merged range
            col_letter = openpyxl.utils.get_column_letter(c)
            cell_coord = f"{col_letter}{r}"
            is_merged = False
            for m_range in ws.merged_cells.ranges:
                if cell_coord in m_range:
                    is_merged = True
                    break
            
            col_w = cols_with_width.get(col_letter, 8.43) # default excel width
            
            # If not merged, and text is significantly longer than column width and wrap_text is not enabled
            if not is_merged and len(val_str) > col_w + 3:
                align = cell.alignment
                is_wrapped = align.wrap_text if align else False
                if not is_wrapped and not val_str.startswith("="):
                    potential_clipping.append((cell_coord, val_str[:30], len(val_str), col_w))
                elif is_wrapped:
                    # Check if row height is enough for wrapped text
                    estimated_lines = (len(val_str) // int(col_w)) + 1
                    needed_height = estimated_lines * 14
                    if r_height and r_height < needed_height:
                        wrap_issues.append((cell_coord, val_str[:30], r_height, needed_height))

    if potential_clipping:
        print(f"  ⚠️ Potential unmerged clipping ({len(potential_clipping)} cells):")
        for coord, text, length, w in potential_clipping[:5]:
            print(f"     {coord}: '{text}...' (len={length}, col_w={w})")
    else:
        print("  ✅ No unmerged text clipping detected.")
        
    if wrap_issues:
        print(f"  ⚠️ Potential row height too small for wrapped text ({len(wrap_issues)} cells):")
        for coord, text, rh, needed in wrap_issues[:5]:
            print(f"     {coord}: '{text}...' (row_h={rh}, needed~={needed})")
    else:
        print("  ✅ Row heights are well-proportioned for wrapped text.")
    print()
