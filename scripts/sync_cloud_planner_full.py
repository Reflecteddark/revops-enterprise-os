"""
Скрипт полной синхронизации RevOps_Founder_Workspace_Planner.xlsx
с облачной Google Таблицей:
https://docs.google.com/spreadsheets/d/1AbaQGHhzar7Qsb9U9TByN57rAmOn36zw4HOD2XVrmTc

Обновляет все ключевые листы и добавляет новый лист '🔥 База_Лидов_HH' (55 горячих лидов).
"""

import sys
from pathlib import Path
import openpyxl
import gspread

sys.stdout.reconfigure(encoding="utf-8")

SPREADSHEET_ID = "1AbaQGHhzar7Qsb9U9TByN57rAmOn36zw4HOD2XVrmTc"
LOCAL_EXCEL = Path(r"C:\Users\strel\Desktop\RevOps_Founder_Workspace_Planner.xlsx")

client = gspread.service_account(filename="service_account.json")
sh = client.open_by_key(SPREADSHEET_ID)
wb = openpyxl.load_workbook(LOCAL_EXCEL, data_only=True)

sheets_to_sync = [
    "🎯 Спринты_и_Задачи",
    "💡 База_Знаний_и_УТП",
    "🤝 Пайплайн_Интеграторов",
    "🎙️ Клиентские_Пилоты",
    "💬 Быстрые_Скрипты",
    "🔥 База_Лидов_HH"
]

for s_name in sheets_to_sync:
    if s_name not in wb.sheetnames:
        continue
    ws_local = wb[s_name]

    # Читаем данные из Excel
    values = []
    for r in range(1, ws_local.max_row + 1):
        row_vals = []
        for c in range(1, ws_local.max_column + 1):
            val = ws_local.cell(r, c).value
            row_vals.append("" if val is None else str(val))
        values.append(row_vals)

    # Проверяем наличие листа в Google Sheets
    try:
        ws_cloud = sh.worksheet(s_name)
    except gspread.WorksheetNotFound:
        ws_cloud = sh.add_worksheet(title=s_name, rows=len(values) + 10, cols=len(values[0]) + 2)
        print(f"[+] Создан новый облачный лист: {s_name}")

    # Записываем значения
    ws_cloud.clear()
    range_str = f"A1:{openpyxl.utils.get_column_letter(len(values[0]))}{len(values)}"
    ws_cloud.update(range_name=range_str, values=values)
    print(f"[✓] Синхронизирован лист '{s_name}' (строк: {len(values)}, колонок: {len(values[0])})")

print("\n🚀 Полная синхронизация с Google Sheets завершена успешно!")
