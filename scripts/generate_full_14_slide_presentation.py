import os
import sys
import subprocess
import shutil
import base64
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

# Read Dmitry Photo
photo_path = os.path.abspath("docs/images/founder-dmitry.jpg")
photo_b64 = ""
if os.path.exists(photo_path):
    with open(photo_path, "rb") as img_f:
        photo_b64 = base64.b64encode(img_f.read()).decode("utf-8")

html_content = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>RevOps Enterprise OS V18.0 — Презентация</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap');

@page {{
    size: 16in 9in;
    margin: 0;
}}

* {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
}}

body {{
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    background: #E2E8F0;
    color: #0F172A;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
}}

.slide {{
    width: 16in;
    height: 9in;
    page-break-after: always;
    position: relative;
    overflow: hidden;
    padding: 55px 80px;
    background: #F1F5F9;
    display: flex;
    flex-direction: column;
}}

.slide.dark-bg {{
    background: #0B1329;
    color: #FFFFFF;
    justify-content: center;
    padding: 70px 100px;
}}

.dark-bg .bg-glow {{
    position: absolute;
    right: -100px;
    bottom: -100px;
    width: 750px;
    height: 750px;
    background: radial-gradient(circle, rgba(37, 99, 235, 0.25) 0%, rgba(15, 23, 42, 0) 70%);
    border-radius: 50%;
    pointer-events: none;
}}

.header-logo {{
    display: flex;
    align-items: center;
    gap: 12px;
    margin-bottom: 35px;
}}

.logo-icon {{
    width: 38px;
    height: 38px;
    background: #2563EB;
    border-radius: 10px;
    display: flex;
    align-items: center;
    justify-content: center;
    color: white;
    font-weight: 800;
    font-size: 20px;
}}

.logo-text {{
    font-size: 26px;
    font-weight: 700;
    letter-spacing: -0.5px;
    color: #FFFFFF;
}}

.slide-header {{
    display: flex;
    justify-content: space-between;
    align-items: flex-start;
    margin-bottom: 30px;
}}

.slide-title {{
    font-size: 36px;
    font-weight: 800;
    color: #0F172A;
    letter-spacing: -0.5px;
    line-height: 1.2;
}}

.slide-subtitle {{
    font-size: 17px;
    color: #64748B;
    margin-top: 5px;
    font-weight: 500;
}}

.corner-logo {{
    display: flex;
    align-items: center;
    gap: 8px;
    opacity: 0.9;
}}

.corner-logo-text {{
    font-size: 18px;
    font-weight: 700;
    color: #2563EB;
}}

.card {{
    background: #FFFFFF;
    border-radius: 20px;
    padding: 28px;
    box-shadow: 0 4px 20px rgba(0, 0, 0, 0.04);
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

/* Slide 1 Styles */
.hero-title {{
    font-size: 60px;
    font-weight: 800;
    line-height: 1.15;
    margin-bottom: 22px;
    letter-spacing: -1px;
    max-width: 1100px;
}}

.hero-subtitle {{
    font-size: 23px;
    color: #94A3B8;
    line-height: 1.5;
    max-width: 950px;
    margin-bottom: 45px;
    font-weight: 400;
}}

.hero-badge {{
    display: inline-flex;
    align-items: center;
    gap: 8px;
    background: rgba(37, 99, 235, 0.2);
    border: 1px solid rgba(59, 130, 246, 0.4);
    padding: 8px 18px;
    border-radius: 30px;
    font-size: 14px;
    font-weight: 600;
    color: #60A5FA;
    margin-bottom: 25px;
    width: fit-content;
}}

/* Grid Helpers */
.quad-grid {{
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 20px;
    flex: 1;
}}

.tri-grid {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 26px;
    flex: 1;
}}

.dual-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 28px;
    flex: 1;
}}

/* Slide 2 */
.top-hub-card {{
    background: #FFFFFF;
    border-radius: 18px;
    padding: 20px 30px;
    text-align: center;
    margin-bottom: 20px;
    border: 1px solid #CBD5E1;
}}

.top-hub-title {{
    font-size: 21px;
    font-weight: 800;
    color: #0F172A;
}}

.top-hub-sub {{
    font-size: 14.5px;
    color: #64748B;
    margin-top: 3px;
}}

.quad-card {{
    background: #FFFFFF;
    border-radius: 18px;
    padding: 24px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

.quad-tag {{
    display: inline-block;
    font-size: 11px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    color: #2563EB;
    background: #EFF6FF;
    padding: 4px 10px;
    border-radius: 6px;
    margin-bottom: 12px;
    width: fit-content;
}}

.quad-title {{
    font-size: 18px;
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 10px;
    line-height: 1.3;
}}

.quad-desc {{
    font-size: 13px;
    color: #64748B;
    line-height: 1.5;
    flex: 1;
}}

.quad-metric {{
    margin-top: 14px;
    padding-top: 12px;
    border-top: 1px solid #F1F5F9;
    font-size: 13px;
    font-weight: 600;
    color: #0F172A;
}}

/* Slide 3 Tri */
.tri-card {{
    background: #FFFFFF;
    border-radius: 22px;
    padding: 40px 32px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

.icon-box {{
    width: 60px;
    height: 60px;
    border-radius: 16px;
    background: #EFF6FF;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 28px;
    margin-bottom: 24px;
    color: #2563EB;
}}

.tri-title {{
    font-size: 21px;
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 14px;
    line-height: 1.3;
}}

.tri-desc {{
    font-size: 15px;
    color: #64748B;
    line-height: 1.55;
}}

/* Slide 4: Comparison vs Traditional Speech Analytics */
.vs-card {{
    border-radius: 22px;
    padding: 35px;
    display: flex;
    flex-direction: column;
}}

.vs-bad {{
    background: #FEF2F2;
    border: 1px solid #FECACA;
}}

.vs-good {{
    background: #ECFDF5;
    border: 2px solid #10B981;
}}

.vs-header {{
    font-size: 22px;
    font-weight: 800;
    margin-bottom: 12px;
}}

.vs-sub {{
    font-size: 14.5px;
    margin-bottom: 22px;
}}

.vs-list {{
    display: flex;
    flex-direction: column;
    gap: 14px;
}}

.vs-item {{
    display: flex;
    align-items: flex-start;
    gap: 12px;
    font-size: 14.5px;
    line-height: 1.45;
}}

/* Audience */
.audience-card {{
    background: #FFFFFF;
    border-radius: 22px;
    padding: 32px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

.audience-header {{
    border-bottom: 2px solid #F1F5F9;
    padding-bottom: 14px;
    margin-bottom: 18px;
}}

.audience-role {{
    font-size: 21px;
    font-weight: 800;
    color: #0F172A;
}}

.audience-sub {{
    font-size: 13.5px;
    color: #64748B;
    margin-top: 3px;
}}

.task-list {{
    display: flex;
    flex-direction: column;
    gap: 12px;
    flex: 1;
}}

.task-item {{
    display: flex;
    align-items: flex-start;
    gap: 12px;
}}

.task-num {{
    width: 24px;
    height: 24px;
    border-radius: 50%;
    background: #EFF6FF;
    color: #2563EB;
    font-weight: 700;
    font-size: 12.5px;
    display: flex;
    align-items: center;
    justify-content: center;
    flex-shrink: 0;
    margin-top: 2px;
}}

.task-text {{
    font-size: 14px;
    color: #334155;
    line-height: 1.4;
}}

/* Steps 1-2-3 */
.steps-grid {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 24px;
    flex: 1;
}}

.step-card {{
    background: #FFFFFF;
    border-radius: 22px;
    padding: 35px 30px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
    position: relative;
    overflow: hidden;
}}

.step-watermark {{
    position: absolute;
    top: -20px;
    right: 20px;
    font-size: 110px;
    font-weight: 900;
    color: #F1F5F9;
    z-index: 1;
    pointer-events: none;
    line-height: 1;
}}

.step-content {{
    position: relative;
    z-index: 2;
    display: flex;
    flex-direction: column;
    flex: 1;
}}

.step-num-badge {{
    width: 40px;
    height: 40px;
    border-radius: 10px;
    background: #2563EB;
    color: #FFFFFF;
    font-weight: 800;
    font-size: 19px;
    display: flex;
    align-items: center;
    justify-content: center;
    margin-bottom: 18px;
}}

.step-title {{
    font-size: 20px;
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 14px;
    line-height: 1.3;
}}

.step-points {{
    list-style: none;
    display: flex;
    flex-direction: column;
    gap: 10px;
}}

.step-point {{
    font-size: 13.5px;
    color: #475569;
    line-height: 1.45;
    position: relative;
    padding-left: 18px;
}}

.step-point::before {{
    content: "•";
    position: absolute;
    left: 4px;
    color: #2563EB;
    font-weight: 900;
    font-size: 16px;
}}

/* Numbers Grid */
.numbers-grid {{
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    grid-template-rows: repeat(2, 1fr);
    gap: 20px;
    flex: 1;
}}

.num-card {{
    background: #FFFFFF;
    border-radius: 18px;
    padding: 26px 20px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
    justify-content: center;
}}

.num-val {{
    font-size: 38px;
    font-weight: 900;
    color: #2563EB;
    letter-spacing: -1px;
    line-height: 1.1;
    margin-bottom: 8px;
}}

.num-label {{
    font-size: 13px;
    color: #475569;
    line-height: 1.4;
}}

/* Facts Grid */
.facts-grid {{
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    grid-template-rows: repeat(2, 1fr);
    gap: 18px;
    flex: 1;
}}

.fact-card {{
    background: #FFFFFF;
    border-radius: 16px;
    padding: 20px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

.check-circle {{
    width: 28px;
    height: 28px;
    border-radius: 50%;
    background: #10B981;
    color: white;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 15px;
    font-weight: 800;
    margin-bottom: 12px;
}}

.fact-title {{
    font-size: 15px;
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 6px;
    line-height: 1.3;
}}

.fact-desc {{
    font-size: 12.5px;
    color: #64748B;
    line-height: 1.4;
}}

/* Principles */
.principle-card {{
    background: #FFFFFF;
    border-radius: 22px;
    padding: 38px 30px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
}}

.principle-title {{
    font-size: 20px;
    font-weight: 800;
    color: #0F172A;
    margin-bottom: 16px;
    line-height: 1.3;
}}

.principle-body {{
    font-size: 14.5px;
    color: #475569;
    line-height: 1.55;
    margin-bottom: 18px;
}}

.principle-highlight {{
    margin-top: auto;
    padding: 12px 16px;
    background: #EFF6FF;
    border-radius: 10px;
    font-size: 13px;
    color: #1E40AF;
    font-weight: 600;
    line-height: 1.4;
}}

/* Roadmap 7 Days */
.roadmap-grid {{
    display: grid;
    grid-template-columns: repeat(5, 1fr);
    gap: 16px;
    flex: 1;
}}

.roadmap-card {{
    background: #FFFFFF;
    border-radius: 18px;
    padding: 24px 20px;
    border: 1px solid #E2E8F0;
    display: flex;
    flex-direction: column;
    position: relative;
}}

.roadmap-day-badge {{
    font-size: 11px;
    font-weight: 800;
    text-transform: uppercase;
    color: #2563EB;
    background: #EFF6FF;
    padding: 4px 10px;
    border-radius: 6px;
    margin-bottom: 12px;
    width: fit-content;
}}

.roadmap-card-title {{
    font-size: 16.5px;
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 10px;
    line-height: 1.3;
}}

.roadmap-card-desc {{
    font-size: 13px;
    color: #64748B;
    line-height: 1.45;
}}

.roadmap-card-result {{
    margin-top: auto;
    padding-top: 12px;
    border-top: 1px solid #F1F5F9;
    font-size: 12.5px;
    font-weight: 600;
    color: #10B981;
}}

/* Slide 14: Final Contact Card + 0 Rubles Offer */
.final-layout {{
    display: grid;
    grid-template-columns: 460px 1fr;
    gap: 35px;
    align-items: center;
    flex: 1;
}}

.founder-card {{
    background: rgba(15, 23, 42, 0.7);
    border: 1px solid rgba(255, 255, 255, 0.12);
    border-radius: 24px;
    padding: 26px;
    display: flex;
    flex-direction: column;
    gap: 16px;
}}

.founder-profile {{
    display: flex;
    align-items: center;
    gap: 16px;
}}

.founder-avatar-box {{
    position: relative;
    width: 68px;
    height: 68px;
    border-radius: 18px;
    overflow: hidden;
    border: 2px solid #3B82F6;
    flex-shrink: 0;
}}

.founder-avatar-img {{
    width: 100%;
    height: 100%;
    object-fit: cover;
}}

.online-dot {{
    position: absolute;
    bottom: -2px;
    right: -2px;
    width: 16px;
    height: 16px;
    border-radius: 50%;
    background: #10B981;
    border: 2px solid #0B1329;
}}

.founder-name {{
    font-size: 20px;
    font-weight: 800;
    color: #FFFFFF;
}}

.founder-title {{
    font-size: 13.5px;
    color: #10B981;
    font-weight: 600;
}}

.founder-tagline {{
    font-size: 12.5px;
    color: #94A3B8;
}}

.founder-bio {{
    font-size: 13px;
    color: #CBD5E1;
    line-height: 1.45;
}}

.contact-btn-row {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 10px;
}}

.btn-contact {{
    background: #1E293B;
    border: 1px solid #334155;
    border-radius: 12px;
    padding: 12px;
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 8px;
    font-size: 13.5px;
    font-weight: 700;
    color: #FFFFFF;
    text-decoration: none;
}}

.btn-tg {{
    background: #0284C7;
    border-color: #38BDF8;
}}

.btn-max {{
    background: #4F46E5;
    border-color: #818CF8;
}}

.bot-btn-row {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 10px;
}}

.bot-card-btn {{
    background: #1E293B;
    border: 1px solid #334155;
    border-radius: 12px;
    padding: 10px 14px;
    display: flex;
    align-items: center;
    gap: 10px;
    text-decoration: none;
}}

.bot-card-title {{
    font-size: 13px;
    font-weight: 700;
    color: #FFFFFF;
}}

.bot-card-sub {{
    font-size: 11px;
    color: #94A3B8;
}}

.email-box {{
    background: #1E293B;
    border: 1px solid #334155;
    border-radius: 12px;
    padding: 10px 16px;
    display: flex;
    align-items: center;
    gap: 12px;
}}

.email-title {{
    font-size: 13.5px;
    font-weight: 700;
    color: #FFFFFF;
}}

.email-sub {{
    font-size: 11px;
    color: #94A3B8;
}}

/* Offer Box */
.offer-box {{
    background: rgba(255, 255, 255, 0.05);
    border: 2px solid #2563EB;
    border-radius: 26px;
    padding: 40px;
    display: flex;
    flex-direction: column;
    justify-content: center;
}}

.offer-badge {{
    display: inline-flex;
    align-items: center;
    gap: 8px;
    background: #10B981;
    color: white;
    font-weight: 800;
    font-size: 13px;
    text-transform: uppercase;
    padding: 6px 16px;
    border-radius: 20px;
    margin-bottom: 20px;
    width: fit-content;
}}

.offer-title {{
    font-size: 34px;
    font-weight: 800;
    color: #FFFFFF;
    line-height: 1.2;
    margin-bottom: 14px;
}}

.offer-desc {{
    font-size: 16px;
    color: #94A3B8;
    line-height: 1.5;
    margin-bottom: 26px;
}}

.offer-steps {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 16px;
    margin-bottom: 30px;
}}

.offer-step-card {{
    background: rgba(15, 23, 42, 0.6);
    border: 1px solid rgba(255, 255, 255, 0.1);
    border-radius: 14px;
    padding: 16px;
}}

.offer-step-num {{
    font-size: 12px;
    font-weight: 700;
    color: #60A5FA;
    margin-bottom: 4px;
}}

.offer-step-text {{
    font-size: 13px;
    color: #E2E8F0;
    line-height: 1.35;
}}

.offer-cta-btn {{
    background: #2563EB;
    color: white;
    padding: 16px 28px;
    border-radius: 14px;
    font-size: 16px;
    font-weight: 800;
    text-align: center;
    text-decoration: none;
    display: inline-block;
    box-shadow: 0 4px 15px rgba(37, 99, 235, 0.4);
}}
</style>
</head>
<body>

<!-- СЛАЙД 1: ОБЛОЖКА -->
<div class="slide dark-bg">
    <div class="bg-glow"></div>
    <div class="header-logo">
        <div class="logo-icon">R</div>
        <div class="logo-text">RevOps Enterprise OS</div>
    </div>
    <div class="hero-badge">ENTERPRISE RELEASE V18.0 • REVOPS ARCHITECTURE</div>
    <h1 class="hero-title">Добро пожаловать в RevOps Enterprise OS!</h1>
    <p class="hero-subtitle">
        Интеллектуальная платформа сквозного контроля выручки, 100% речевого ИИ-аудита диалогов и оперативного управления отделом продаж для B2B-компаний.
    </p>
    <div style="font-size: 16px; color: #64748B; font-weight: 600;">
        Готовое решение для Собственников (CEO) и Руководителей отделов продаж (РОП)
    </div>
</div>

<!-- СЛАЙД 2: АРХИТЕКТУРА ЭКОСИСТЕМЫ -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Архитектура экосистемы RevOps Core</div>
            <div class="slide-subtitle">Единая технологическая среда для полной прозрачности коммерческого блока</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="top-hub-card">
        <div class="top-hub-title">RevOps Core Platform (Единый Источник Правды — SSOT)</div>
        <div class="top-hub-sub">Центральное вычислительное ядро и дашборд, агрегирующий сделки, звонки, финансы и KPI</div>
    </div>

    <div class="quad-grid">
        <div class="quad-card">
            <div class="quad-tag">ИИ-Модуль</div>
            <div class="quad-title">AI Speech Inspector</div>
            <div class="quad-desc">
                Собственный сервер речевой аналитики. Автоматически транскрибирует 100% звонков менеджеров и оценивает их по 13 жестким стандартам продаж.
            </div>
            <div class="quad-metric">100% звонков без человека</div>
        </div>

        <div class="quad-card">
            <div class="quad-tag">CRM-Модуль</div>
            <div class="quad-title">amoCRM / B24 Reverse ETL</div>
            <div class="quad-desc">
                Двусторонний бесшовный коннектор. За 15 секунд возвращает в карточку сделки текстовый аудит, оценку, теги и задачи Next Step.
            </div>
            <div class="quad-metric">Запись в сделку за 15 сек</div>
        </div>

        <div class="quad-card">
            <div class="quad-tag">Менеджмент</div>
            <div class="quad-title">Пульт РОПа & Воронка SLA</div>
            <div class="quad-desc">
                Рабочий стол руководителя: 15-минутная утренняя планерка, контроль зависших КП (>48ч), платежный календарь дебиторки (DSO).
            </div>
            <div class="quad-metric">Планерка за 15 минут</div>
        </div>

        <div class="quad-card">
            <div class="quad-tag">Security & Bot</div>
            <div class="quad-title">Telegram War Room Bot</div>
            <div class="quad-desc">
                Мгновенный радар критических инцидентов. Оповещает РОПа в течение 15 минут о сливе крупной сделки, формирует утренние и вечерние сводки.
            </div>
            <div class="quad-metric">Алерт при срыве сделки</div>
        </div>
    </div>
</div>

<!-- СЛАЙД 3: ЧЕМ ЗАНИМАЕТСЯ REVOPS OS -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Чем занимается RevOps Enterprise OS?</div>
            <div class="slide-subtitle">Мы превращаем непредсказуемый хаос в продажах в точную математическую систему</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="tri-grid">
        <div class="tri-card">
            <div class="icon-box">🛡️</div>
            <div class="tri-title">Оцифровывает и ликвидирует финансовые потери</div>
            <div class="tri-desc">
                Находит 4 главные системные утечки отдела продаж (слив на звонках, зависание КП >48ч, брошенные отказники, задержка оплат) и возвращает до 28% упущенной выручки компании.
            </div>
        </div>

        <div class="tri-card">
            <div class="icon-box">🎙️</div>
            <div class="tri-title">Слушает и объективно оценивает 100% звонков</div>
            <div class="tri-desc">
                Полностью освобождает РОПа от рутинного прослушивания сотен аудиозаписей. Нейросеть беспристрастно аудирует каждый диалог и дает персональные Coaching Tips менеджерам.
            </div>
        </div>

        <div class="tri-card">
            <div class="icon-box">📊</div>
            <div class="tri-title">Дает руководству управляемость и прогноз кассы</div>
            <div class="tri-desc">
                Заменяет субъективные отчеты на единый пульт с прогнозом закрытия месяца (Run-Rate), контролем регламентов SLA и прозрачной формулой динамической мотивации (Payroll).
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 4: ПОЧЕМУ ОБЫЧНАЯ РЕЧЕВКА НЕ СПАСАЕТ (НОВЫЙ СЛАЙД) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Почему обычная речевая аналитика не спасает продажи?</div>
            <div class="slide-subtitle">Разница между «просто расшифровкой текста» и боевой операционной системой управления выручкой</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="dual-grid">
        <div class="vs-card vs-bad">
            <div class="vs-header" style="color: #DC2626;">❌ Обычная речевая аналитика из телефонии</div>
            <div class="vs-sub" style="color: #7F1D1D;">Пассивный архив текста, который никто не читает</div>
            <div class="vs-list">
                <div class="vs-item">
                    <span>⚠️</span>
                    <div><strong>Просто полотно расшифрованного текста:</strong> РОПу всё равно нужно тратить по 3 часа в день, чтобы вчитываться в сотни диалогов.</div>
                </div>
                <div class="vs-item">
                    <span>⚠️</span>
                    <div><strong>Бесполезные «облака тегов»:</strong> считает слова-паразиты («здравствуйте», «спасибо»), но не понимает, закрыта ли сделка.</div>
                </div>
                <div class="vs-item">
                    <span>⚠️</span>
                    <div><strong>Оторвана от денег и воронки:</strong> не знает сумму сделки, зависание КП и дату оплаты выставленного счета.</div>
                </div>
                <div class="vs-item">
                    <span>⚠️</span>
                    <div><strong>Нет связи с CRM:</strong> не возвращает комментарии, не ставит задачи менеджерам и не влияет на зарплату.</div>
                </div>
            </div>
        </div>

        <div class="vs-card vs-good">
            <div class="vs-header" style="color: #059669;">✅ Боевая платформа RevOps Enterprise OS</div>
            <div class="vs-sub" style="color: #064E3B;">Операционная система, которая возвращает деньги в кассу</div>
            <div class="vs-list">
                <div class="vs-item">
                    <span>🎯</span>
                    <div><strong>Оценка по 13 критериям продаж:</strong> проверяет квалификацию ЛПР, бюджет, дедлайны и железный Next Step с датой и временем.</div>
                </div>
                <div class="vs-item">
                    <span>⚡</span>
                    <div><strong>Развернутое саммари в CRM за 15 сек:</strong> менеджер кладет трубку — в карточке сделки уже лежит готовый разбор и советы РОПа.</div>
                </div>
                <div class="vs-item">
                    <span>🛡️</span>
                    <div><strong>Автопостановка задач:</strong> забыл назначить следующий контакт? Робот сам ставит жесткую задачу менеджеру до конца дня.</div>
                </div>
                <div class="vs-item">
                    <span>💰</span>
                    <div><strong>Прямая связь с выручкой:</strong> Пульт РОПа за 15 минут в день + автоматический дисконт бонуса менеджера за брак речи.</div>
                </div>
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 5: КТО НАШИ КЛИЕНТЫ -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Кто наши клиенты и какие задачи мы решаем?</div>
            <div class="slide-subtitle">Специализированное решение для двух ключевых лидеров коммерческого блока</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="dual-grid">
        <div class="audience-card">
            <div class="audience-header">
                <div class="audience-role">Собственники и Генеральные директора (CEO)</div>
                <div class="audience-sub">Компании сегмента B2B, оптовой торговли, производства и услуг</div>
            </div>
            <div class="task-list">
                <div class="task-item"><div class="task-num">1</div><div class="task-text">Прекращаем слив рекламного бюджета на неквалифицированных звонках менеджеров.</div></div>
                <div class="task-item"><div class="task-num">2</div><div class="task-text">Обеспечиваем математически точный прогноз кассы к концу месяца (Run-Rate).</div></div>
                <div class="task-item"><div class="task-num">3</div><div class="task-text">Оцифровываем скрытые финансовые потери коммерческого блока в рублях.</div></div>
                <div class="task-item"><div class="task-num">4</div><div class="task-text">Исключаем зависимость бизнеса от «незаменимых» и токсичных менеджеров-звезд.</div></div>
                <div class="task-item"><div class="task-num">5</div><div class="task-text">Внедряем честный расчет зарплат: выплата бонусов только за соблюдение регламентов.</div></div>
                <div class="task-item"><div class="task-num">6</div><div class="task-text">Защищаем компанию от кассовых разрывов через контроль дебиторской задолженности (DSO).</div></div>
            </div>
        </div>

        <div class="audience-card">
            <div class="audience-header">
                <div class="audience-role">Руководители отделов продаж (РОП / Коммерческий директор)</div>
                <div class="audience-sub">Управление командами от 3 до 50+ продавцов в CRM</div>
            </div>
            <div class="task-list">
                <div class="task-item"><div class="task-num">1</div><div class="task-text">Сокращаем утреннюю планерку с 1.5 часов до 15 минут строго по фактам и цифрам.</div></div>
                <div class="task-item"><div class="task-num">2</div><div class="task-text">Экономим до 3 часов в день за счет автоматического 100% аудита звонков нейросетью.</div></div>
                <div class="task-item"><div class="task-num">3</div><div class="task-text">Мгновенно узнаем о срыве крупных сделок и успеваем перехватить клиента за 15 минут.</div></div>
                <div class="task-item"><div class="task-num">4</div><div class="task-text">Гарантируем 100% контроль железобетонного Next Step (дата и время контакта).</div></div>
                <div class="task-item"><div class="task-num">5</div><div class="task-text">Получаем готовые Coaching Tips — персональные фразы-подсказки для каждого менеджера.</div></div>
                <div class="task-item"><div class="task-num">6</div><div class="task-text">Устраняем рутину ручного контроля заполнения карточек и задач в amoCRM.</div></div>
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 6: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 1 - АУДИТ ЗВОНКОВ) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Как это работает: Речевой ИИ-Аудит Звонков</div>
            <div class="slide-subtitle">Технология объективного контроля 100% телефонных переговоров без участия человека</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="steps-grid">
        <div class="step-card">
            <div class="step-watermark">1</div>
            <div class="step-content">
                <div class="step-num-badge">1</div>
                <div class="step-title">Бесшовный захват аудио из CRM</div>
                <ul class="step-points">
                    <li class="step-point">Интеграция перехватывает аудиозапись из любой телефонии (Манго, Телфин, Sipuni, Мои Звонки).</li>
                    <li class="step-point">Захватываются звонки 100% менеджеров по всей компании.</li>
                    <li class="step-point">0 ручных действий: менеджер просто говорит по телефону.</li>
                    <li class="step-point">Шифрованная передача аудиофайла в защищенный контур.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">2</div>
            <div class="step-content">
                <div class="step-num-badge">2</div>
                <div class="step-title">Скоринг по 13 критериям продаж</div>
                <ul class="step-points">
                    <li class="step-point">ИИ транскрибирует речь и сопоставляет диалог с жестким регламентом компании.</li>
                    <li class="step-point">Бинарная проверка (0/1): приветствие, ЛПР, боли, бюджет, дедлайны КП.</li>
                    <li class="step-point">Железобетонный контроль: зафиксирована ли точная дата и время следующего шага.</li>
                    <li class="step-point">Формирование итогового балла качества от 0 до 100.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">3</div>
            <div class="step-content">
                <div class="step-num-badge">3</div>
                <div class="step-title">Формирование Coaching Tip</div>
                <ul class="step-points">
                    <li class="step-point">Определение зоны риска: красный (&lt;70), желтый (70-84) или зеленый (85+).</li>
                    <li class="step-point">Генерация конкретной рекомендации: какую фразу нужно было сказать вместо ошибки.</li>
                    <li class="step-point">Анализ эмоционального тона диалога и детекция раздражения клиента.</li>
                    <li class="step-point">Фиксация данных в единой аналитической базе.</li>
                </ul>
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 7: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 2 - AMOCRM) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Как это работает: Двусторонняя Автоматизация amoCRM</div>
            <div class="slide-subtitle">Робот возвращает аналитику, задачи и теги обратно в CRM за 15 секунд</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="steps-grid">
        <div class="step-card">
            <div class="step-watermark">1</div>
            <div class="step-content">
                <div class="step-num-badge">1</div>
                <div class="step-title">Развернутое саммари в сделку</div>
                <ul class="step-points">
                    <li class="step-point">Через 15 секунд после звонка в ленте карточки появляется структурированное примечание.</li>
                    <li class="step-point">Краткий пересказ сути договоренностей (чтение за 20 секунд).</li>
                    <li class="step-point">Баллы по 13 критериям и совет РОПа.</li>
                    <li class="step-point">Заполнение числового поля «Оценка ИИ (0–100)» для фильтрации воронки.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">2</div>
            <div class="step-content">
                <div class="step-num-badge">2</div>
                <div class="step-title">Автопостановка задач Next Step</div>
                <ul class="step-points">
                    <li class="step-point">Если менеджер забыл назначить дату перезвона, ИИ сам ставит жесткую задачу до конца дня.</li>
                    <li class="step-point">Если клиент попросил КП к четвергу в 14:00, робот создаст задачу с точным дедлайном.</li>
                    <li class="step-point">100% защита от «забытых» и потерянных клиентов.</li>
                    <li class="step-point">Авто-заполнение полей бюджета и потребностей из разговора.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">3</div>
            <div class="step-content">
                <div class="step-num-badge">3</div>
                <div class="step-title">Умное тегирование и Follow-Up</div>
                <ul class="step-points">
                    <li class="step-point">Авто-теги: #Брак_Речи, #Слив_Клиента, #Эталонный_Звонок, #Возражение_Дорого.</li>
                    <li class="step-point">Генерация готового текста сообщения в WhatsApp/Telegram для отправки клиенту.</li>
                    <li class="step-point">Менеджеру остается нажать 1 кнопку для отправки резюме встречи.</li>
                    <li class="step-point">Мгновенная сегментация воронки по реальным возражениям.</li>
                </ul>
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 8: КАК ЭТО РАБОТАЕТ (МОДУЛЬ 3 - УПРАВЛЕНИЕ И БЕЗОПАСНОСТЬ) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Как это работает: Операционный Пульт и Безопасность</div>
            <div class="slide-subtitle">Инструменты оперативного контроля выручки и соответствие стандартам безопасности РФ</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="steps-grid">
        <div class="step-card">
            <div class="step-watermark">1</div>
            <div class="step-content">
                <div class="step-num-badge">1</div>
                <div class="step-title">Утренний Пульт РОПа (15 минут)</div>
                <ul class="step-points">
                    <li class="step-point">Оценка темпа месяца (Pacing): идем ли по графику плана выручки.</li>
                    <li class="step-point">Светофор воронки: автоматическое выявление зависших сделок на этапе КП (&gt;48ч).</li>
                    <li class="step-point">ТОП-5 критических рисков выручки: список сделок, требующих немедленного дожима.</li>
                    <li class="step-point">Ежедневная постановка фокусов команде за 15 минут.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">2</div>
            <div class="step-content">
                <div class="step-num-badge">2</div>
                <div class="step-title">Telegram War Room & Алерты</div>
                <ul class="step-points">
                    <li class="step-point">Красная кнопка РОПу: при срыве сделки с чеком &gt;300 000 ₽ — сигнал в Telegram за 15 минут.</li>
                    <li class="step-point">РОП успевает лично перезвонить клиенту и спасти контракт.</li>
                    <li class="step-point">Утренняя сводка в 09:00: рейтинг менеджеров и средний балл речи отдела.</li>
                    <li class="step-point">Вечерний отчет собственнику в 19:00: касса за день и прогноз месяца.</li>
                </ul>
            </div>
        </div>

        <div class="step-card">
            <div class="step-watermark">3</div>
            <div class="step-content">
                <div class="step-num-badge">3</div>
                <div class="step-title">Безопасность и 152-ФЗ РФ</div>
                <ul class="step-points">
                    <li class="step-point">Полное соблюдение Федерального закона 152-ФЗ: сервера расположены на территории РФ.</li>
                    <li class="step-point">Изоляция данных Multi-Tenant: каждый клиент работает в закрытом контуре Tenant UUID.</li>
                    <li class="step-point">Защита IP и коммерческой тайны: закрытое ядро формул и промптов.</li>
                    <li class="step-point">Шифрование трафика и разграничение прав доступа.</li>
                </ul>
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 9: ВНЕДРЕНИЕ БЕЗ САБОТАЖА И УВОЛЬНЕНИЙ (НОВЫЙ СЛАЙД) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Внедрение без саботажа и увольнений сотрудников</div>
            <div class="slide-subtitle">Как искусственный интеллект становится союзником менеджера, а не «надзирателем»</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="tri-grid">
        <div class="principle-card">
            <div class="principle-title">1. Личный тренер, а не надзиратель</div>
            <div class="principle-body">
                Разборы ошибок и рекомендации приходят менеджеру в личные сообщения сразу после звонка, а не вывешиваются на всеобщее обозрение. Сотрудник видит свои точки роста тет-а-тет, без стресса и токсичности.
            </div>
            <div class="principle-highlight">
                РОП превращается из надзирателя в сильного ментора и наставника команды.
            </div>
        </div>

        <div class="principle-card">
            <div class="principle-title">2. Снятие 80% рутины с сейлзов</div>
            <div class="principle-body">
                Продавцы ненавидят заполнять CRM. RevOps OS сам пишет краткое саммари диалога, ставит напоминания о договоренностях и формирует готовый Follow-Up текст для клиента в WhatsApp.
            </div>
            <div class="principle-highlight">
                Менеджеры освобождают до 1.5 часов в день на реальное общение с клиентами.
            </div>
        </div>

        <div class="principle-card">
            <div class="principle-title">3. Честные и прозрачные бонусы</div>
            <div class="principle-body">
                Исключаются любимчики и споры о зарплате в конце месяца. Система рассчитывает премию строго по формуле: выполнение плана + соблюдение стандартов. Кто продает по правилам — зарабатывает больше всех.
            </div>
            <div class="principle-highlight">
                Прозрачная мотивация устраняет текучку сильных кадров в отделе продаж.
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 10: ДОРОЖНАЯ КАРТА ПИЛОТА (НОВЫЙ СЛАЙД) -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Дорожная карта пилота: от 0 до найденных денег за 7 дней</div>
            <div class="slide-subtitle">Пошаговый план быстрого старта без остановки текущих продаж</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="roadmap-grid">
        <div class="roadmap-card">
            <div class="roadmap-day-badge">День 1 • 15 минут</div>
            <div class="roadmap-card-title">Подключение CRM и Телефонии</div>
            <div class="roadmap-card-desc">
                Выдача гостевого доступа к amoCRM / Битрикс24. Активация защищенного вебхука и запуск перехвата звонков.
            </div>
            <div class="roadmap-card-result">✓ Коннектор активен</div>
        </div>

        <div class="roadmap-card">
            <div class="roadmap-day-badge">Дни 2–3 • 24 часа</div>
            <div class="roadmap-card-title">Калибровка 13 Стандартов</div>
            <div class="roadmap-card-desc">
                Нейросеть слушает первые 50–100 звонков. Адаптация критериев оценки под специфику вашей B2B-ниши и чеков.
            </div>
            <div class="roadmap-card-result">✓ Промпт откалиброван</div>
        </div>

        <div class="roadmap-card">
            <div class="roadmap-day-badge">Дни 4–5 • Контроль</div>
            <div class="roadmap-card-title">Автопостановка Задач и Теги</div>
            <div class="roadmap-card-desc">
                Включение умного тегирования (#Брак_Речи, #Слив_Клиента) и автопостановки жестких задач Next Step в карточках сделок.
            </div>
            <div class="roadmap-card-result">✓ Сливы ликвидированы</div>
        </div>

        <div class="roadmap-card">
            <div class="roadmap-day-badge">День 6 • Операционка</div>
            <div class="roadmap-card-title">Запуск 15-мин Пульта РОПа</div>
            <div class="roadmap-card-desc">
                Проведение первой утренней планерки по новому регламенту. Разбор ТОП-5 критических зависших сделок на этапе КП.
            </div>
            <div class="roadmap-card-result">✓ Планерка за 15 минут</div>
        </div>

        <div class="roadmap-card">
            <div class="roadmap-day-badge">День 7 • Деньги</div>
            <div class="roadmap-card-title">Финансовый Отчет Собственнику</div>
            <div class="roadmap-card-desc">
                Демонстрация объема предотвращенных финансовых потерь, пересчет Run-Rate кассы и расчет ROI пилота.
            </div>
            <div class="roadmap-card-result">✓ Окупаемость 3x–5x</div>
        </div>
    </div>
</div>

<!-- СЛАЙД 11: КОМПАНИЯ В ЦИФРАХ -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Платформа RevOps OS в цифрах</div>
            <div class="slide-subtitle">Ключевые измеримые метрики надежности и финансовой эффективности системы</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="numbers-grid">
        <div class="num-card"><div class="num-val">100%</div><div class="num-label">звонков аудируются искусственным интеллектом без участия человека</div></div>
        <div class="num-card"><div class="num-val">15 минут</div><div class="num-label">длительность утренней планерки РОПа по фактам и цифрам вместо 1.5 часов</div></div>
        <div class="num-card"><div class="num-val">+18–34%</div><div class="num-label">средний прирост чистой выручки клиентов в первые 60 дней внедрения</div></div>
        <div class="num-card"><div class="num-val">&lt; 48 часов</div><div class="num-label">жесткий норматив нахождения сделки на этапе КП под контролем SLA</div></div>
        <div class="num-card"><div class="num-val">13 стандартов</div><div class="num-label">объективной оценки диалогов в эталонной матрице речевого аудита</div></div>
        <div class="num-card"><div class="num-val">118 тестов</div><div class="num-label">автоматической проверки целостности математики в модуле QA Suite</div></div>
        <div class="num-card"><div class="num-val">4 утечки</div><div class="num-label">финансовых потерь ликвидируются (Next Step, КП, отказники, дебиторка)</div></div>
        <div class="num-card"><div class="num-val">15 минут</div><div class="num-label">техническое время полного развертывания платформы для нового клиента</div></div>
    </div>
</div>

<!-- СЛАЙД 12: ФАКТЫ О REVOPS OS -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Факты о RevOps Enterprise OS</div>
            <div class="slide-subtitle">Технологические и юридические стандарты промышленного решения</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="facts-grid">
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Единый контур SSOT</div><div class="fact-desc">Единственная система, объединяющая речевой аудит, CRM-автоматизацию и пульт РОПа в один экран.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Двусторонний обмен CRM</div><div class="fact-desc">Прямая интеграция с amoCRM и Битрикс24 без установки сторонних расширений в браузеры сотрудников.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">100% 152-ФЗ РФ Compliant</div><div class="fact-desc">Серверная инфраструктура и базы данных размещены в аккредитованных дата-центрах на территории РФ.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Закрытый AI Black Box</div><div class="fact-desc">Промпты и скоринговые модели защищены на выделенном сервере и постоянно дообучаются.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Быстрый старт за 1 день</div><div class="fact-desc">Запуск пилота без остановки работы отдела продаж. Первые данные доступны уже на следующий день.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Адаптация под B2B-нишу</div><div class="fact-desc">Калибровка 13 стандартов речи под специфику опта, сложных услуг, дистрибуции или производства.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Динамический Payroll</div><div class="fact-desc">Автоматический расчет премий с дисконтом 30% за систематический брак переговоров.</div></div>
        <div class="fact-card"><div class="check-circle">✓</div><div class="fact-title">Полная окупаемость пилота</div><div class="fact-desc">Возврат инвестиций в 3–5 раз за счет спасения всего 1–2 зависших сделок.</div></div>
    </div>
</div>

<!-- СЛАЙД 13: НАШИ ПРИНЦИПЫ -->
<div class="slide">
    <div class="slide-header">
        <div>
            <div class="slide-title">Почему мы достигаем таких результатов?</div>
            <div class="slide-subtitle">Три фундаментальных принципа архитектуры RevOps Enterprise OS</div>
        </div>
        <div class="corner-logo"><span class="corner-logo-text">RevOps OS</span></div>
    </div>

    <div class="tri-grid">
        <div class="principle-card">
            <div class="principle-title">1. Цифры и факты вместо субъективных отчетов</div>
            <div class="principle-body">
                Менеджеры больше не могут прикрывать неудачи фразами «клиент думает» или «он не берет трубку». Руководство видит объективную стенограмму разговора, соблюдение этапов и точные таймкоды ключевых договоренностей.
            </div>
            <div class="principle-highlight">
                Управленческие решения принимаются на базе реальных данных, а не догадок.
            </div>
        </div>

        <div class="principle-card">
            <div class="principle-title">2. Фокус на живой выручке, а не звонках ради галочки</div>
            <div class="principle-body">
                Мы оцениваем не бессмысленную длительность звонков, а продвижение сделки к деньгам: квалификацию ЛПР, выявление бюджета, сокращение времени согласования КП и соблюдение сроков оплаты счетов (DSO).
            </div>
            <div class="principle-highlight">
                Каждый диалог менеджера измеряется его влиянием на итоговую кассу месяца.
            </div>
        </div>

        <div class="principle-card">
            <div class="principle-title">3. Обучение команды без стресса и микроменеджмента</div>
            <div class="principle-body">
                Искусственный интеллект подсказывает ошибки тактично и сразу после звонка через персональные Coaching Tips. Менеджеры растут в квалификации самостоятельно, а РОП становится наставником, а не контролером.
            </div>
            <div class="principle-highlight">
                Здоровая атмосфера в отделе продаж и рост конверсий без текучки кадров.
            </div>
        </div>
    </div>
</div>

<!-- СЛАЙД 14: КОНТАКТЫ, ДМИТРИЙ ФЕДОТОВ, ОФФЕР 0 РУБЛЕЙ (ФИНАЛЬНЫЙ) -->
<div class="slide dark-bg">
    <div class="bg-glow"></div>
    <div class="final-layout">
        
        <!-- Левая колонка: Карточка Дмитрия Федотова из референса -->
        <div class="founder-card">
            <div class="founder-profile">
                <div class="founder-avatar-box">
                    <img src="data:image/jpeg;base64,{photo_b64}" alt="Дмитрий Федотов" class="founder-avatar-img">
                    <div class="online-dot"></div>
                </div>
                <div>
                    <div class="founder-name">Дмитрий Федотов</div>
                    <div class="founder-title">Основатель ai-rop.ru</div>
                    <div class="founder-tagline">B2B Sales & RevOps</div>
                </div>
            </div>

            <div class="founder-bio">
                Лично курирую каждый пилотный проект и отвечаю на вопросы по интеграции.
            </div>

            <!-- Кнопки Telegram и MAX -->
            <div class="contact-btn-row">
                <a href="https://t.me/dm1918" class="btn-contact btn-tg">
                    ✈️ @dm1918
                </a>
                <a href="https://max.ru/u/f9LHodD0cOLRcNKRakz94FpZjmUJ25lYzJUMzPBJyLKwxM1Tzzai1aF-dTg" class="btn-contact btn-max">
                    💬 MAX
                </a>
            </div>

            <!-- Кнопки Ботов -->
            <div class="bot-btn-row">
                <a href="https://max.ru/se14526668_bot" class="bot-card-btn">
                    <span style="font-size: 20px;">💬</span>
                    <div>
                        <div class="bot-card-title">MAX Бот</div>
                        <div class="bot-card-sub">Экспресс-аудит...</div>
                    </div>
                </a>
                <a href="https://t.me/RevOps_Super_Audit_Bot?start=audit3" class="bot-card-btn">
                    <span style="font-size: 20px;">🤖</span>
                    <div>
                        <div class="bot-card-title">TG-Бот</div>
                        <div class="bot-card-sub">Скоринг звонков...</div>
                    </div>
                </a>
            </div>

            <!-- Email -->
            <div class="email-box">
                <span style="font-size: 20px;">✉️</span>
                <div>
                    <div class="email-title">info@ai-rop.ru</div>
                    <div class="email-sub">Официальная почта для юрлиц и КП</div>
                </div>
            </div>
        </div>

        <!-- Правая колонка: Оффер «Аудит 3 звонков за 0 рублей» -->
        <div class="offer-box">
            <div class="offer-badge">🎁 Входной Тест-Драйв Без Оплаты</div>
            <div class="offer-title">Экспресс-Аудит 3 Звонков = 0 Рублей</div>
            <div class="offer-desc">
                Хотите увидеть, как нейросеть разберет реальные переговоры ваших продавцов? Загрузите 3 аудиозаписи в нашего Telegram или MAX-бота прямо сейчас — вы получите глубокий отчет за 2 минуты абсолютно бесплатно!
            </div>

            <div class="offer-steps">
                <div class="offer-step-card">
                    <div class="offer-step-num">ШАГ 1</div>
                    <div class="offer-step-text">Откройте бота в Telegram или MAX по ссылкам слева.</div>
                </div>
                <div class="offer-step-card">
                    <div class="offer-step-num">ШАГ 2</div>
                    <div class="offer-step-text">Перешлите 3 аудиофайла любых звонков вашей команды.</div>
                </div>
                <div class="offer-step-card">
                    <div class="offer-step-num">ШАГ 3</div>
                    <div class="offer-step-text">Получите PDF-отчет с баллами по 13 критериям и советами РОПа.</div>
                </div>
            </div>

            <div style="display: flex; gap: 20px; align-items: center;">
                <a href="https://t.me/RevOps_Super_Audit_Bot?start=audit3" class="offer-cta-btn">
                    🚀 Запустить бесплатный аудит 3 звонков в Telegram
                </a>
                <span style="font-size: 13.5px; color: #94A3B8;">Сайт: <strong>ai-rop.ru</strong></span>
            </div>
        </div>

    </div>
</div>

</body>
</html>
"""

# Save HTML
html_path = os.path.abspath("docs/Презентация_RevOps_Enterprise_OS_14_Слайдов.html")
with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_content)
print(f"HTML presentation written to: {html_path}")

# Render PDF via Edge
edge_bin = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
pdf_path = os.path.abspath("docs/Презентация_RevOps_Enterprise_OS_14_Слайдов.pdf")
pres_pdf = os.path.abspath("presentation/Презентация_RevOps_Enterprise_OS_14_Слайдов.pdf")

desktop_dir = r"C:\Users\strel\Desktop\RevOps Platform\Презентация\Преза на отправку"
desktop_pdf = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_14_Слайдов.pdf")
desktop_pdf_std = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_11_Слайдов.pdf") # overwrite old one too

cmd = [
    edge_bin,
    "--headless",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={pdf_path}",
    html_path
]

res = subprocess.run(cmd, capture_output=True, text=True)
print("Edge return code:", res.returncode)

if os.path.exists(pdf_path):
    shutil.copyfile(pdf_path, pres_pdf)
    shutil.copyfile(pdf_path, desktop_pdf)
    shutil.copyfile(pdf_path, desktop_pdf_std)
    print(f"SUCCESS: Generated 14-Slide PDF: {pdf_path} ({os.path.getsize(pdf_path):,} bytes)")
    print(f"SUCCESS: Copied to desktop: {desktop_pdf}")
else:
    print("FAILED to generate PDF via Edge")

# -------------------------------------------------------------
# Generate Matching Landscape DOCX (14 Slides)
# -------------------------------------------------------------
doc = docx.Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

header = section.header
p_head = header.paragraphs[0]
p_head.text = "REVOPS ENTERPRISE OS V18.0  |  ОФИЦИАЛЬНАЯ ПРЕЗЕНТАЦИЯ (14 СЛАЙДОВ)"
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_head.runs[0].font.size = Pt(8)
p_head.runs[0].font.name = "Calibri"

footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.text = "RevOps Architecture  •  ai-rop.ru  •  Дмитрий Федотов (@dm1918)  •  Аудит 3 звонков: 0 руб."
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_foot.runs[0].font.size = Pt(8)
p_foot.runs[0].font.name = "Calibri"

def add_slide_docx(doc, title, slide_num, content_func):
    if slide_num > 1:
        doc.add_page_break()
    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(6)
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run(f"СЛАЙД {slide_num} ИЗ 14")
    r_num.font.name = "Calibri"
    r_num.font.size = Pt(9)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(37, 99, 235)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(title)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(17)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    content_func(doc)

# Slide 1
def s1(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Интеллектуальная платформа сквозного контроля выручки, 100% речевого ИИ-аудита диалогов и оперативного управления отделом продаж для B2B-компаний.\n\nГотовое промышленное решение для Собственников (CEO) и Руководителей отделов продаж (РОП).")
    r.font.size = Pt(11)

# Slide 2
def s2(doc):
    p = doc.add_paragraph()
    p.add_run("Единая технологическая среда для полной прозрачности коммерческого блока.\n\n"
              "• RevOps Core Platform (SSOT) — центральное ядро, объединяющее звонки, сделки и финансы.\n"
              "• AI Speech Inspector — транскрибация и оценка 100% звонков по 13 жестким стандартам.\n"
              "• amoCRM / B24 Reverse ETL — двусторонний обмен: запись саммари, тегов и задач за 15 секунд.\n"
              "• Пульт РОПа & Воронка SLA — утренняя планерка за 15 минут, контроль зависших КП (>48ч).\n"
              "• Telegram War Room Bot — радар критических инцидентов и мгновенные алерты руководству.")
    p.runs[0].font.size = Pt(10)

# Slide 3
def s3(doc):
    p = doc.add_paragraph()
    p.add_run("Три фундаментальные задачи платформы:\n\n"
              "1. Оцифровывает и ликвидирует финансовые потери: устраняет 4 утечки (сливы звонков, зависание КП, брошенные отказники, дебиторку) и возвращает до 28% выручки.\n\n"
              "2. Слушает и объективно оценивает 100% звонков: освобождает РОПа от рутинного прослушивания сотен записей, дает готовые Coaching Tips.\n\n"
              "3. Дает руководству управляемость и прогноз кассы: точный прогноз закрытия месяца (Run-Rate) и прозрачная динамическая мотивация (Payroll).")
    p.runs[0].font.size = Pt(10)

# Slide 4: Comparison vs Traditional
def s4(doc):
    p = doc.add_paragraph()
    p.add_run("❌ ОБЫЧНАЯ РЕЧЕВАЯ АНАЛИТИКА ТЕЛЕФОНИИ:\n"
              "• Просто стенограмма текста: РОПу нужно читать часы диалогов.\n"
              "• Облака тегов и слов-паразитов: не понимает, закрыта ли сделка.\n"
              "• Оторвана от денег и воронки: не знает сумму сделки и стадию оплаты.\n"
              "• Нет связи с CRM: не ставит задачи и не меняет зарплату менеджера.\n\n"
              "✅ БОЕВАЯ ПЛАТФОРМА REVOPS ENTERPRISE OS:\n"
              "• Оценка по 13 критериям продаж (ЛПР, бюджет, дедлайны, Next Step).\n"
              "• Развернутое саммари в карточку сделки за 15 секунд после звонка.\n"
              "• Автопостановка жестких задач Next Step при потере контакта.\n"
              "• Прямая связь с кассой: Пульт РОПа за 15 минут + дисконт бонуса за брак речи.")
    p.runs[0].font.size = Pt(10)

# Slide 5: Audience
def s5(doc):
    p = doc.add_paragraph()
    p.add_run("👑 ЗАДАЧИ СОБСТВЕННИКОВ (CEO):\n"
              "1. Прекращаем слив рекламного бюджета на неквалифицированных звонках.\n"
              "2. Обеспечиваем математически точный прогноз кассы к концу месяца (Run-Rate).\n"
              "3. Оцифровываем скрытые финансовые потери коммерческого блока в рублях.\n"
              "4. Исключаем зависимость от токсичных «звезд».\n"
              "5. Внедряем честный расчет бонусов за соблюдение регламентов.\n"
              "6. Защищаем компанию от кассовых разрывов через контроль дебиторки (DSO).\n\n"
              "📋 ЗАДАЧИ РУКОВОДИТЕЛЕЙ ПРОДАЖ (РОП):\n"
              "1. Сокращаем планерку с 1.5 часов до 15 минут строго по фактам.\n"
              "2. Экономим 3 часа в день на прослушивании аудиозаписей.\n"
              "3. Мгновенно узнаем о срыве крупных сделок и успеваем спасти клиента за 15 минут.\n"
              "4. Гарантируем 100% контроль даты и времени следующего контакта (Next Step).\n"
              "5. Получаем готовые подсказки Coaching Tips для прокачки менеджеров.\n"
              "6. Устраняем рутину ручного контроля карточек в amoCRM.")
    p.runs[0].font.size = Pt(9.5)

# Slide 6: Module 1
def s6(doc):
    p = doc.add_paragraph()
    p.add_run("• Шаг 1. Бесшовный захват аудио: перехват из Манго, Телфина, Sipuni, Моих Звонков без участия людей.\n"
              "• Шаг 2. Скоринг по 13 критериям: проверка ЛПР, потребностей, бюджета, дедлайнов и фиксации даты следующего шага.\n"
              "• Шаг 3. Формирование Coaching Tip: балл 0-100, выделение ошибок и готовая фраза-подсказка.")
    p.runs[0].font.size = Pt(10)

# Slide 7: Module 2
def s7(doc):
    p = doc.add_paragraph()
    p.add_run("• Шаг 1. Развернутое саммари в сделку за 15 секунд + заполнение числового поля «Оценка ИИ (0–100)».\n"
              "• Шаг 2. Автопостановка задач Next Step при потере контакта или по обещанию клиенту к точному времени.\n"
              "• Шаг 3. Умное тегирование (#Брак_Речи, #Возражение_Дорого) и генерация Follow-Up текста для WhatsApp.")
    p.runs[0].font.size = Pt(10)

# Slide 8: Module 3
def s8(doc):
    p = doc.add_paragraph()
    p.add_run("• 1. 15-минутный Пульт РОПа: темп месяца (Pacing), воронка AS-IS, ТОП-5 критических рисков выручки.\n"
              "• 2. Telegram War Room: красная тревога РОПу за 15 минут при срыве сделки >300 000 ₽.\n"
              "• 3. Безопасность и 152-ФЗ РФ: сервера в РФ, изоляция данных по Tenant UUID, закрытый AI Black Box.")
    p.runs[0].font.size = Pt(10)

# Slide 9: No sabotage
def s9(doc):
    p = doc.add_paragraph()
    p.add_run("1. Личный тренер, а не надзиратель: ошибки и советы приходят менеджеру тет-а-тет в личку, без публичного позора. РОП становится наставником.\n\n"
              "2. Снятие 80% рутины: система сама пишет саммари, ставит напоминания и генерирует текст для WhatsApp. Менеджеры освобождают 1.5 часа в день на продажи.\n\n"
              "3. Прозрачные деньги и справедливые бонусы: премия считается строго по формуле. Кто продает по правилам — зарабатывает больше всех. Исключается текучка кадров.")
    p.runs[0].font.size = Pt(10)

# Slide 10: Roadmap 7 Days
def s10(doc):
    p = doc.add_paragraph()
    p.add_run("• День 1 (15 минут): Подключение amoCRM / телефонии, запуск защищенного вебхука.\n"
              "• Дни 2–3 (24 часа): Захват первых 100 звонков, калибровка 13 стандартов под B2B-нишу.\n"
              "• Дни 4–5 (Контроль): Активация автопостановки задач Next Step и умных тегов в сделках.\n"
              "• День 6 (Операционка): Первая планерка по Пульту РОПа, разбор ТОП-5 рисковых зависших сделок.\n"
              "• День 7 (Деньги): Финансовый отчет Собственнику, объем предотвращенных потерь, окупаемость пилота 3x–5x.")
    p.runs[0].font.size = Pt(10)

# Slide 11: Numbers
def s11(doc):
    p = doc.add_paragraph()
    p.add_run("• 100% звонков аудируются ИИ без участия человека\n"
              "• 15 минут — длительность ежедневной планерки РОПа по фактам\n"
              "• +18–34% средний рост чистой выручки в первые 60 дней\n"
              "• < 48 часов — норматив нахождения сделки на этапе КП под контролем SLA\n"
              "• 13 стандартов в эталонной речевой матрице\n"
              "• 118 автоматических тестов математики в QA Suite\n"
              "• 4 финансовые утечки ликвидируются под ключ\n"
              "• 15 минут — техническое время развертывания нового клиента")
    p.runs[0].font.size = Pt(10)

# Slide 12: Facts
def s12(doc):
    p = doc.add_paragraph()
    p.add_run("✓ Единый контур SSOT (речевой аудит + CRM + Пульт РОПа в одном месте).\n"
              "✓ Двусторонний обмен CRM без установки сторонних виджетов в браузер.\n"
              "✓ 100% 152-ФЗ РФ Compliant (сервера и базы на территории России).\n"
              "✓ Закрытый AI Black Box (промпты защищены и постоянно дообучаются).\n"
              "✓ Быстрый старт за 1 день без остановки продаж.\n"
              "✓ Адаптация под B2B-нишу (опт, сложные услуги, дистрибуция, производство).\n"
              "✓ Динамический Payroll (авторасчет бонусов с дисконтом 30% за брак речи).\n"
              "✓ Окупаемость пилота в 3–5 раз за счет спасения всего 1–2 зависших сделок.")
    p.runs[0].font.size = Pt(10)

# Slide 13: Principles
def s13(doc):
    p = doc.add_paragraph()
    p.add_run("1. Цифры и факты вместо субъективных отчетов: исключаем отговорки «клиент думает» — опираемся на стенограмму, таймкоды и зафиксированные дедлайны.\n\n"
              "2. Фокус на живой выручке, а не звонках ради галочки: измеряем продвижение к деньгам и сокращение цикла закрытия сделок.\n\n"
              "3. Обучение команды без стресса и микроменеджмента: ИИ подсказывает ошибки тактично и сразу, здоровая атмосфера в коллективе.")
    p.runs[0].font.size = Pt(10)

# Slide 14: Founder Contacts & Offer 0 Rubles
def s14(doc):
    p = doc.add_paragraph()
    p.add_run("👤 ДМИТРИЙ ФЕДОТОВ — Основатель ai-rop.ru (B2B Sales & RevOps)\n"
              "«Лично курирую каждый пилотный проект и отвечаю на вопросы по интеграции».\n\n"
              "• Telegram: @dm1918 (https://t.me/dm1918)\n"
              "• MAX Профиль: https://max.ru/u/f9LHodD0cOLRcNKRakz94FpZjmUJ25lYzJUMzPBJyLKwxM1Tzzai1aF-dTg\n"
              "• MAX Бот (Экспресс-аудит): https://max.ru/se14526668_bot\n"
              "• Telegram Бот (Скоринг звонков): https://t.me/RevOps_Super_Audit_Bot?start=audit3\n"
              "• Официальная почта для юрлиц и КП: info@ai-rop.ru\n"
              "• Официальный сайт: https://ai-rop.ru\n\n"
              "🎁 СПЕЦИАЛЬНЫЙ ВХОДНОЙ ОФФЕР: ЭКСПРЕСС-АУДИТ 3 ЗВОНКОВ = 0 РУБЛЕЙ\n"
              "Загрузите 3 аудиозаписи вашей команды в бота (Telegram или MAX) прямо сейчас — и получите глубокий аудит по 13 критериям за 2 минуты абсолютно бесплатно!")
    p.runs[0].font.size = Pt(10.5)

add_slide_docx(doc, "Добро пожаловать в RevOps Enterprise OS!", 1, s1)
add_slide_docx(doc, "Архитектура экосистемы RevOps Core", 2, s2)
add_slide_docx(doc, "Чем занимается RevOps Enterprise OS?", 3, s3)
add_slide_docx(doc, "Почему обычная речевая аналитика не спасает продажи?", 4, s4)
add_slide_docx(doc, "Кто наши клиенты и какие задачи мы решаем?", 5, s5)
add_slide_docx(doc, "Как это работает: Модуль 1. Речевой ИИ-Аудит Звонков", 6, s6)
add_slide_docx(doc, "Как это работает: Модуль 2. Двусторонняя Автоматизация amoCRM", 7, s7)
add_slide_docx(doc, "Как это работает: Модуль 3. Операционный Пульт РОПа и Безопасность", 8, s8)
add_slide_docx(doc, "Внедрение без саботажа и увольнений сотрудников", 9, s9)
add_slide_docx(doc, "Дорожная карта пилота: от 0 до найденных денег за 7 дней", 10, s10)
add_slide_docx(doc, "Платформа RevOps OS в цифрах", 11, s11)
add_slide_docx(doc, "Ключевые факты о RevOps Enterprise OS", 12, s12)
add_slide_docx(doc, "Почему мы достигаем таких результатов? Наши принципы", 13, s13)
add_slide_docx(doc, "Контакты и Бесплатный Тест-Драйв: Аудит 3 Звонков = 0 ₽", 14, s14)

docx_path = os.path.abspath("docs/Презентация_RevOps_Enterprise_OS_14_Слайдов.docx")
doc.save(docx_path)
print(f"Word presentation written to: {docx_path}")

pres_docx = os.path.abspath("presentation/Презентация_RevOps_Enterprise_OS_14_Слайдов.docx")
shutil.copyfile(docx_path, pres_docx)

desktop_docx = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_14_Слайдов.docx")
desktop_docx_std = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_11_Слайдов.docx") # overwrite old one too
shutil.copyfile(docx_path, desktop_docx)
shutil.copyfile(docx_path, desktop_docx_std)
print(f"SUCCESS: Copied DOCX to desktop: {desktop_docx}")

html_dest = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_14_Слайдов.html")
html_dest_std = os.path.join(desktop_dir, "Презентация_RevOps_Enterprise_OS_11_Слайдов.html")
shutil.copyfile(html_path, html_dest)
shutil.copyfile(html_path, html_dest_std)
print(f"SUCCESS: Copied HTML to desktop: {html_dest}")

print("ALL 14 SLIDES ASSEMBLED AND DEPLOYED.")
