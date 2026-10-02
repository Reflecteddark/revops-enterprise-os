import datetime
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_invoice_docx(out_path: Path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    p_header = doc.add_paragraph()
    r = p_header.add_run("Внимание! Оплата данного счета означает согласие с условиями публичного Договора-оферты на сайте ai-rop.ru/offer.html\n"
                         "Уведомление об оплате обязательно направлять в Telegram: @dm1918")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.italic = True
    r.font.color.rgb = RGBColor(100, 116, 139)

    # Bank details table
    b_table = doc.add_table(rows=4, cols=4)
    b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(2.5), Inches(1.2), Inches(1.1), Inches(2.2)]

    b_data = [
        ("АО «ТИНЬКОФФ БАНК»\nг. Москва\nБанк получателя", "БИК", "044525974", ""),
        ("", "Сч. №", "30101810145250000974", ""),
        ("ИНН 772485920194\nФедотов Дмитрий (RevOps OS)\nПолучатель", "Сч. №", "40802810500003849120", ""),
        ("", "", "", "")
    ]
    # Simple bank table
    for r_idx in range(4):
        for c_idx in range(4):
            c = b_table.cell(r_idx, c_idx)
            c.width = widths[c_idx]

    b_table.cell(0, 0).paragraphs[0].text = "АО «ТИНЬКОФФ БАНК» г. Москва\nБанк получателя"
    b_table.cell(0, 1).paragraphs[0].text = "БИК\nСч. №"
    b_table.cell(0, 2).paragraphs[0].text = "044525974\n30101810145250000974"
    b_table.cell(2, 0).paragraphs[0].text = "ИНН 772485920194\nФедотов Дмитрий (RevOps OS / ai-rop.ru)\nПолучатель"
    b_table.cell(2, 1).paragraphs[0].text = "Сч. №"
    b_table.cell(2, 2).paragraphs[0].text = "40802810500003849120"

    p_inv = doc.add_paragraph()
    p_inv.paragraph_format.space_before = Pt(12)
    p_inv.paragraph_format.space_after = Pt(6)
    r_inv = p_inv.add_run("Счет на оплату № PLT-01 от «____» ____________ 2026 г.")
    r_inv.font.name = "Arial"
    r_inv.font.size = Pt(13)
    r_inv.font.bold = True

    p_parties = doc.add_paragraph()
    p_parties.paragraph_format.space_after = Pt(10)
    p_parties.paragraph_format.line_spacing = 1.2
    r_p = p_parties.add_run(
        "Поставщик (Исполнитель): ИП / Плательщик НПД Федотов Дмитрий, ИНН 772485920194, р/с 40802810500003849120 в АО «ТИНЬКОФФ БАНК», БИК 044525974, тел/TG: @dm1918\n"
        "Покупатель (Заказчик): ____________________________________________________________________, ИНН ____________________, КПП ____________________"
    )
    r_p.font.name = "Arial"
    r_p.font.size = Pt(9)

    # Items table
    i_table = doc.add_table(rows=2, cols=6)
    i_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["№", "Товары (работы, услуги)", "Кол-во", "Ед.", "Цена", "Сумма"]
    row_h = i_table.rows[0]
    for idx, name in enumerate(headers):
        cell = row_h.cells[idx]
        set_cell_shading(cell, "1E2E4A")
        p = cell.paragraphs[0]
        run = p.add_run(name)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    item_row = i_table.rows[1]
    vals = [
        "1",
        "Услуги по проведению 7-дневного пилотного спринта внедрения программного комплекса ИИ-супервизии звонков «RevOps OS» по Договору-оферте от 02.10.2026 (со 100% гарантией возврата)",
        "1",
        "усл.",
        "14 900,00",
        "14 900,00"
    ]
    for idx, val in enumerate(vals):
        cell = item_row.cells[idx]
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        if idx in (0, 2, 3):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif idx in (4, 5):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p_total = doc.add_paragraph()
    p_total.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_total.paragraph_format.space_before = Pt(6)
    r_t = p_total.add_run("Итого к оплате: 14 900,00 ₽\nВ том числе НДС: Без НДС (УСН / НПД)")
    r_t.font.name = "Arial"
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)

    p_words = doc.add_paragraph()
    p_words.paragraph_format.space_before = Pt(4)
    r_w = p_words.add_run("Всего наименований 1, на сумму 14 900,00 ₽\nЧетырнадцать тысяч девятьсот рублей 00 копеек")
    r_w.font.name = "Arial"
    r_w.font.size = Pt(9)
    r_w.font.bold = True

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(25)
    r_s = p_sig.add_run("Исполнитель: ____________________ / Федотов Д. /\t\tМ.П.")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(9.5)

    doc.save(str(out_path))
    print(f"✓ Invoice DOCX generated: {out_path}")


def generate_act_docx(out_path: Path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    p_act = doc.add_paragraph()
    p_act.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_act.paragraph_format.space_after = Pt(4)
    r_act = p_act.add_run("АКТ СДАЧИ-ПРИЕМКИ ОКАЗАННЫХ УСЛУГ № PLT-01\nпо Договору-оферте от 02.10.2026 г.")
    r_act.font.name = "Arial"
    r_act.font.size = Pt(12)
    r_act.font.bold = True

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_after = Pt(10)
    r_d = p_date.add_run("г. Москва / Дистанционно\t\t\t\t\t«____» ____________ 2026 г.")
    r_d.font.name = "Arial"
    r_d.font.size = Pt(9)
    r_d.font.color.rgb = RGBColor(100, 116, 139)

    p_parties = doc.add_paragraph()
    p_parties.paragraph_format.space_after = Pt(10)
    p_parties.paragraph_format.line_spacing = 1.15
    r_p = p_parties.add_run(
        "Исполнитель: ИП / Плательщик НПД Федотов Дмитрий (ИНН: 772485920194), с одной стороны, и\n"
        "Заказчик: ____________________________________________________________________ (ИНН: ____________________), с другой стороны,\n"
        "составили настоящий Акт о том, что Исполнитель оказал, а Заказчик принял следующие услуги:"
    )
    r_p.font.name = "Arial"
    r_p.font.size = Pt(9)

    # Table
    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["№", "Наименование работ (услуг)", "Кол-во", "Ед.", "Цена", "Сумма"]
    row_h = table.rows[0]
    for idx, name in enumerate(headers):
        cell = row_h.cells[idx]
        set_cell_shading(cell, "1E2E4A")
        p = cell.paragraphs[0]
        run = p.add_run(name)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    item_row = table.rows[1]
    vals = [
        "1",
        "Комплекс услуг по 7-дневному пилотному внедрению программного сервиса ИИ-супервизии звонков «RevOps OS» (подключение CRM, 100% аудит звонков по 13 критериям, Telegram-дайджесты и PDF-карта упущенной выручки)",
        "1",
        "усл.",
        "14 900,00",
        "14 900,00"
    ]
    for idx, val in enumerate(vals):
        cell = item_row.cells[idx]
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        if idx in (0, 2, 3):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif idx in (4, 5):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p_total = doc.add_paragraph()
    p_total.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_total.paragraph_format.space_before = Pt(6)
    r_t = p_total.add_run("Итого: 14 900,00 ₽\nБез НДС (УСН / НПД)")
    r_t.font.name = "Arial"
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)

    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(6)
    p_concl.paragraph_format.space_after = Pt(16)
    r_c = p_concl.add_run(
        "Вышеперечисленные услуги оказаны в полном объеме, своевременно и надлежащего качества. "
        "Стороны взаимных финансовых и материальных претензий друг к другу не имеют."
    )
    r_c.font.name = "Arial"
    r_c.font.size = Pt(9)

    p_sigs = doc.add_paragraph()
    r_s = p_sigs.add_run(
        "ОТ ИСПОЛНИТЕЛЯ:\t\t\t\t\tОТ ЗАКАЗЧИКА:\n\n"
        "_________________ / Федотов Д. /\t\t\t_________________ / ______________ /\n"
        "М.П.\t\t\t\t\t\t\tМ.П."
    )
    r_s.font.name = "Arial"
    r_s.font.size = Pt(9)

    doc.save(str(out_path))
    print(f"✓ Act DOCX generated: {out_path}")


def main():
    docs_dir = Path("docs")
    docs_dir.mkdir(exist_ok=True)
    generate_invoice_docx(docs_dir / "Бланк_Счета_на_оплату_Пилота_14900.docx")
    generate_act_docx(docs_dir / "Типовой_Акт_Оказанных_Услуг.docx")

if __name__ == "__main__":
    main()
