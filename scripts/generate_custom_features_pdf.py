import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import subprocess

sys.stdout.reconfigure(encoding="utf-8")

# Palette RevOps Brand
COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate Navy #0F172A
COLOR_ACCENT = RGBColor(37, 99, 235)      # Royal Blue #2563EB
COLOR_SECONDARY = RGBColor(71, 85, 105)   # Muted Slate #475569
COLOR_SUCCESS = RGBColor(16, 185, 129)    # Emerald Green #10B981
COLOR_WARNING = RGBColor(245, 158, 11)    # Amber Gold #F59E0B
COLOR_DANGER = RGBColor(239, 68, 68)      # Ruby Red #EF4444
COLOR_MUTED = RGBColor(100, 116, 139)     # Gray #64748B

HEX_BG_LIGHT = "F8FAFC"
HEX_BG_HEADER = "0F172A"
HEX_BORDER = "E2E8F0"
HEX_ACCENT_BG = "EFF6FF"
HEX_SUCCESS_BG = "ECFDF5"
HEX_WARNING_BG = "FFFBEB"
HEX_DANGER_BG = "FEF2F2"

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header
        header = section.header
        p_head = header.paragraphs[0]
        p_head.text = "REVOPS ENTERPRISE  |  КАТАЛОГ КАСТОМНЫХ ОПЦИЙ ИНТЕГРАЦИИ AMOCRM"
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.runs[0].font.size = Pt(8)
        p_head.runs[0].font.color.rgb = COLOR_MUTED
        p_head.runs[0].font.name = "Calibri"
        
        # Footer
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.text = "RevOps AI Integration Engine  •  Двусторонний обмен amoCRM  •  Конфиденциально для Клиента"
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.runs[0].font.size = Pt(8)
        p_foot.runs[0].font.color.rgb = COLOR_MUTED
        p_foot.runs[0].font.name = "Calibri"

def add_callout(doc, text, title="ВАЖНО", alert_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    if alert_type == "info":
        bg_color, border_color, title_color = HEX_ACCENT_BG, "2563EB", COLOR_ACCENT
    elif alert_type == "success":
        bg_color, border_color, title_color = HEX_SUCCESS_BG, "10B981", COLOR_SUCCESS
    elif alert_type == "warning":
        bg_color, border_color, title_color = HEX_WARNING_BG, "F59E0B", COLOR_WARNING
    else:
        bg_color, border_color, title_color = HEX_DANGER_BG, "EF4444", COLOR_DANGER
    
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="28" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = title_color
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_PRIMARY
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_ACCENT

def add_p(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_SECONDARY
    return p

def add_feature_box(doc, code, title, description, business_value, default_state=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    set_cell_shading(cell, HEX_BG_LIGHT)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    border_color = "2563EB" if default_state else "94A3B8"
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="{border_color}"/><w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    chk = "[ x ]" if default_state else "[   ]"
    r_chk = p.add_run(f"{chk} ОПЦИЯ {code}: {title}\n")
    r_chk.bold = True
    r_chk.font.name = "Calibri"
    r_chk.font.size = Pt(10.5)
    r_chk.font.color.rgb = COLOR_PRIMARY
    
    r_desc = p.add_run(f"Как работает: {description}\n")
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = COLOR_SECONDARY
    
    r_val = p.add_run(f"💡 Бизнес-эффект: {business_value}")
    r_val.font.name = "Calibri"
    r_val.font.size = Pt(9)
    r_val.font.bold = True
    r_val.font.color.rgb = COLOR_SUCCESS if default_state else COLOR_ACCENT
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def format_custom_table(tbl, col_widths, headers, rows_data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_shading(hdr_cells[i], HEX_BG_HEADER)
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    
    for row_idx, r_data in enumerate(rows_data):
        row = tbl.add_row()
        cells = row.cells
        bg_col = HEX_BG_LIGHT if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cells[c_idx].text = str(val)
            set_cell_shading(cells[c_idx], bg_col)
            set_cell_margins(cells[c_idx], top=70, bottom=70, left=100, right=100)
            p = cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if len(p.runs) > 0:
                p.runs[0].font.name = "Calibri"
                p.runs[0].font.size = Pt(8.5)
                p.runs[0].font.color.rgb = COLOR_PRIMARY
    
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def generate_custom_features_catalog():
    doc = docx.Document()
    add_header_footer(doc)
    
    # Title Block
    p_badge = doc.add_paragraph()
    p_badge.paragraph_format.space_before = Pt(8)
    p_badge.paragraph_format.space_after = Pt(4)
    r_badge = p_badge.add_run("ENTERPRISE INTEGRATION BLUEPRINT  •  AMOCRM REVERSE ETL")
    r_badge.font.name = "Calibri"
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = COLOR_ACCENT
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("Каталог Кастомных Возможностей\nИнтеграции RevOps AI + amoCRM")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Полное меню функциональных модулей двустороннего обмена. Выберите опции галочками для персонализации вашей CRM-системы под ключ в рамках единого подключения.")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    add_callout(doc, 
        "Архитектурное преимущество: Все перечисленные ниже функции реализуются через ОДНУ штатную интеграцию в амоМаркете. "
        "Интеграция работает как двусторонний мост: она автоматически забирает 100% звонков и сделок всей компании, "
        "а затем возвращает в amoCRM аналитику, задачи, теги и рекомендации без задержек (15–20 секунд после звонка).",
        title="ЕДИНАЯ ТОЧКА ПОДКЛЮЧЕНИЯ", alert_type="info")
    
    # -------------------------------------------------------------
    # МОДУЛЬ 1: ОБРАТНАЯ СВЯЗЬ В КАРТОЧКЕ СДЕЛКИ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 1. Мгновенная Обратная Связь в Карточке Сделки")
    
    add_feature_box(doc, "1.1", "Развернутое примечание-аудит в ленту сделки",
        "Через 15 секунд после завершения разговора нейросеть формирует структурированное примечание прямо в таймлайне сделки: "
        "краткое саммари диалога (о чем говорили), статус выполнения 13 стандартов (0/1) и конкретную рекомендацию от РОПа.",
        "РОПу не нужно слушать 20 минут аудиозаписи — он за 15 секунд чтения понимает контекст и качество переговоров.",
        default_state=True)
    
    add_feature_box(doc, "1.2", "Числовое поле «Оценка ИИ (0–100)» в карточке сделки",
        "Система создает в amoCRM кастомное поле и автоматически записывает туда балл качества последнего диалога. "
        "Позволяет настроить быстрый фильтр в канбане: например, подсветить все сделки с оценкой ниже 70.",
        "Управленческий рентген: РОП одним кликом видит сделки, где клиенты недовольны или менеджеры ошибаются.",
        default_state=True)
    
    add_feature_box(doc, "1.3", "Кликабельные таймкоды ключевых моментов звонка",
        "В текст примечания вставляются метки времени: [01:12] Озвучивание цены, [02:30] Возражение по срокам, [04:10] Фиксация шага. "
        "При клике аудио перематывается на нужный момент.",
        "Экономия 90% времени РОПа при выборочном контроле сложных переговоров.",
        default_state=False)
    
    add_feature_box(doc, "1.4", "Умная автоматическая расстановка тегов",
        "ИИ автоматически вешает теги на сделку в зависимости от хода диалога:\n"
        "• #Брак_Речи или #Слив_Клиента — при оценке < 70;\n"
        "• #Эталонный_Звонок — при оценке ≥ 90 (для обучения стажеров);\n"
        "• Смысловые теги: #Возражение_Дорого, #Выбирает_Конкурента, #Требует_Скидку, #Ждет_Спецификацию.",
        "Мгновенная сегментация воронки по реальным возражениям клиентов без ручной разметки менеджерами.",
        default_state=True)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 2: АВТОЗАПОЛНЕНИЕ ПОЛЕЙ И СНЯТИЕ РУТИНЫ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 2. Умное Автозаполнение Полей CRM из Разговора")
    
    add_feature_box(doc, "2.1", "Распознавание и авто-ввод бюджета сделки",
        "Если клиент в разговоре озвучивает сумму («рассчитываем уложиться в 750 тысяч рублей»), нейросеть вычленяет цифру "
        "и автоматически заполняет стандартное поле «Бюджет сделки» в amoCRM.",
        "Повышение точности прогноза выручки: менеджеры часто ленятся вносить бюджет руками.",
        default_state=True)
    
    add_feature_box(doc, "2.2", "Автозаполнение поля «Потребности / Боли клиента»",
        "ИИ анализирует ответы клиента на вопросы менеджера и формирует тезисный список ключевых пожеланий: "
        "«Нужна поставка партии до 25 октября, критичен сертификат ГОСТ, оплата 50/50».",
        "Исключение потери важных технических деталей перед формированием счета или договора.",
        default_state=True)
    
    add_feature_box(doc, "2.3", "Авто-фиксация названий упомянутых конкурентов",
        "Если клиент произносит: «Мы сейчас также запросили коммерческое у компании Х», название компании-конкурента "
        "автоматически заносится в отдельное поле «Конкурент».",
        "Маркетинг получает честную статистику: с кем чаще всего сталкивается отдел продаж.",
        default_state=False)
    
    add_feature_box(doc, "2.4", "Распознавание реквизитов юрлица (ИНН / КПП) из звонка",
        "Если клиент диктует ИНН компании, ИИ фиксирует его в карточке контакта/компании для бухгалтерии.",
        "Ускорение выставления первичных счетов на 1–2 часа.",
        default_state=False)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 3: АВТОПОСТАНОВКА ЗАДАЧ И КОНТРОЛЬ ДИСЦИПЛИНЫ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 3. Автопостановка Задач и Контроль Дисциплины")
    
    add_feature_box(doc, "3.1", "Автопостановка задачи при потере Next Step (Ликвидатор сливов)",
        "Если менеджер завершил разговор, не назначив точную дату и время следующего контакта, робот мгновенно ставит задачу "
        "этому менеджеру: «⚠️ СРОЧНО связаться с клиентом и зафиксировать следующий шаг» со сроком исполнения до конца дня.",
        "Полная ликвидация «забытых клиентов» — 0% сделок остаются брошенными в воздухе.",
        default_state=True)
    
    add_feature_box(doc, "3.2", "Автоматическое создание задач по обещаниям клиента",
        "Если в звонке клиент сказал: «Пришлите обновленный прайс в среду в 15:00», система ставит задачу с дедлайном "
        "ровно на среду 15:00: «Отправить обновленный прайс (по договоренности в звонке)».",
        "Менеджеры больше не держат договоренности в блокнотах или голове — CRM сама ведет сделку.",
        default_state=True)
    
    add_feature_box(doc, "3.3", "Контроль скорости первого ответа (First Response Time SLA)",
        "Система замеряет время от падения нового лида с сайта/рекламы до первого исходящего звонка менеджера. "
        "Если прошло более 15 минут, а звонка нет — задача автоматически эскалируется РОПу.",
        "Рост конверсии из нового лида на +15-20% за счет мгновенного первого контакта.",
        default_state=False)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 4: КОНТРОЛЬ ЭТИКИ, СТОП-СЛОВ И НАСТРОЕНИЯ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 4. Контроль Этики, Стоп-Слов и Анализ Тональности")
    
    add_feature_box(doc, "4.1", "Детектор токсичности, раздражения и риска конфликта",
        "ИИ анализирует лексику и эмоциональный фон диалога. При обнаружении скрытого конфликта, грубости менеджера "
        "или раздражения клиента на сделку вешается тег #Риск_Конфликта, а аудио уходит РОПу.",
        "Защита репутации бренда и спасение недовольных VIP-клиентов до того, как они напишут негативный отзыв.",
        default_state=True)
    
    add_feature_box(doc, "4.2", "Мониторинг запрещенных фраз и стоп-слов компании",
        "Контроль соблюдения Tone of Voice компании. Фиксация паразитных фраз («я не в курсе», «это не ко мне», «смотрите на сайте»), "
        "а также несогласованных обещаний («я сделаю вам максимальную скидку в обход правил»).",
        "Гарантия соблюдения стандартов компании и предотвращение демпинга менеджерами.",
        default_state=False)
    
    add_feature_box(doc, "4.3", "Анализ соотношения времени речи (Talk-to-Listen Ratio)",
        "Замеряет процент времени, который говорил менеджер vs клиент. Если менеджер говорит 80% времени, "
        "система выдает подсказку: «Менеджер вел монолог вместо выявления потребностей». Идеальный бенчмарк: 45% менеджер / 55% клиент.",
        "Обучение сейлзов искусству активного слушания и задавания правильных вопросов.",
        default_state=False)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 5: ГЕНЕРАЦИЯ FOLLOW-UP СООБЩЕНИЙ В МЕССЕНДЖЕРЫ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 5. Авто-Генерация Follow-Up Сообщений (WhatsApp / Telegram)")
    
    add_feature_box(doc, "5.1", "Генерация текста резюме встречи для отправки клиенту",
        "Сразу после звонка ИИ генерирует готовое вежливое сообщение для клиента в WhatsApp/Telegram:\n"
        "«Иван Иванович, добрый день! Спасибо за звонок. Договорились: 1) До среды высылаем КП по модели Х; 2) Созваниваемся в четверг в 14:00 для обсуждения деталей. С уважением, Алексей».",
        "Менеджеру остается только скопировать текст в 1 клик и отправить клиенту. Клиент чувствует высший уровень сервиса.",
        default_state=True)
    
    add_feature_box(doc, "5.2", "Автоматическая отправка Follow-Up через CRM-чат (Wazzup / Radist)",
        "Если в amoCRM подключен WhatsApp-интегратор, сгенерированное сообщение может отправляться клиенту автоматически "
        "через 3 минуты после успешного разговора.",
        "100% фиксация договоренностей в письменном виде без траты рабочего времени продавца.",
        default_state=False)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 6: МГНОВЕННЫЕ TELEGRAM-АЛЕРТЫ РУКОВОДСТВУ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 6. Мгновенные Telegram-Алерты Руководству")
    
    add_feature_box(doc, "6.1", "Красная тревога РОПу: срыв сделки с чеком выше порогового",
        "Если сумма сделки превышает заданный порог (например, 300 000 ₽), а оценка звонка упала ниже 70 баллов, "
        "РОПу в Telegram прилетает мгновенный сигнал со ссылкой на карточку, аудиозаписью и списком ошибок.",
        "Возможность перезвонить клиенту в течение 15 минут со словами: «Здравствуйте, я руководитель отдела, хочу лично проконтролировать ваш заказ» и спасти крупную сделку.",
        default_state=True)
    
    add_feature_box(doc, "6.2", "Утренняя командная сводка РОПу (в 09:00)",
        "Ежедневная карточка в Telegram-чат руководства: рейтинг менеджеров за вчера, общий средний балл речи по отделу, "
        "количество звонков без Next Step и список зависших счетов.",
        "РОП начинает день с полной ясности по отделу еще до того, как открыл ноутбук.",
        default_state=True)
    
    add_feature_box(doc, "6.3", "Вечерний финансовый отчет Собственнику (в 19:00)",
        "Сводка для генерального директора: факт поступлений за день, выполнение плана в %, прогноз закрытия месяца "
        "и объем предотвращенных потерь.",
        "Собственник держит руку на пульсе бизнеса в 1 сообщении без необходимости заходить в CRM.",
        default_state=True)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 7: АВТОМАТИЗАЦИЯ ВОРОНКИ И СМЕНЫ СТАТУСОВ
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 7. Умное Движение Сделок по Воронке (Auto-Pipeline)")
    
    add_feature_box(doc, "7.1", "Авто-перевод сделки на этап «Счет выставлен»",
        "Если клиент в разговоре подтвердил выставление счета («да, выставляйте, реквизиты у вас на почте»), "
        "система автоматически перемещает сделку на этап выставленного счета или создает задачу на бухгалтера.",
        "Сокращение цикла сделки на 1–2 рабочих дня за счет мгновенной передачи в биллинг.",
        default_state=False)
    
    add_feature_box(doc, "7.2", "Интеллектуальная фиксация реальной причины отказа",
        "Если клиент категорически отказывается от продолжения работы, ИИ вычленяет истинную причину из разговора "
        "(«дорого», «купили у конкурента Х», «закрыли проект») и записывает её в аналитику перед переводом в отказ.",
        "Устранение ложных причин («клиент не выходит на связь»), которыми менеджеры прикрывают свои ошибки.",
        default_state=True)
    
    # -------------------------------------------------------------
    # МОДУЛЬ 8: ПЕРСОНАЛЬНЫЙ AI-КОУЧ МЕНЕДЖЕРА
    # -------------------------------------------------------------
    add_heading_1(doc, "Модуль 8. Персональный AI-Тренер Продавца (Геймификация)")
    
    add_feature_box(doc, "8.1", "Личные подсказки менеджеру в Telegram сразу после звонка",
        "Менеджер получает в личные сообщения бота тактичный разбор: «Алексей, отличный диалог (85 баллов)! "
        "Ты круто отработал возражение по цене, но забыл уточнить, кто еще принимает решение. В следующий раз попробуй спросить так:...».",
        "Сотрудник учится непрерывно на каждом звонке. Снижение нагрузки на РОПа по индивидуальному обучению.",
        default_state=True)
    
    add_feature_box(doc, "8.2", "Недельный дашборд личного прогресса сотрудника",
        "Каждую пятницу менеджер видит свой тренд: динамику роста баллов, количество звонков без ошибок, "
        "личный вклад в предотвращенные потери и прогноз своего бонуса за качество.",
        "Позитивная мотивация и профессиональный азарт внутри команды отдела продаж.",
        default_state=False)
    
    # -------------------------------------------------------------
    # БЛАНК СОГЛАСОВАНИЯ КОНФИГУРАЦИИ
    # -------------------------------------------------------------
    add_heading_1(doc, "Бланк Согласования Конфигурации Интеграции")
    add_p(doc, "Заполните таблицу ниже или отметьте необходимые номера опций при согласовании технического задания:")
    
    summary_tbl = doc.add_table(rows=1, cols=4)
    format_custom_table(summary_tbl, [1.0, 2.5, 1.8, 1.7],
        ["Код", "Функциональный модуль", "Рекомендация RevOps", "Выбор Клиента"],
        [
            ["Модуль 1", "Обратная связь в карточке сделки (1.1–1.4)", "🟢 Базовый стандарт", "[  ] Включить полностью"],
            ["Модуль 2", "Автозаполнение полей CRM (2.1–2.4)", "🟢 Рекомендуется", "[  ] Поля бюджета и болей"],
            ["Модуль 3", "Автопостановка задач Next Step (3.1–3.3)", "🟢 Критично для продаж", "[  ] Включить полностью"],
            ["Модуль 4", "Контроль этики и стоп-слов (4.1–4.3)", "🟡 Опционально", "[  ] Только конфликтные"],
            ["Модуль 5", "Follow-Up сообщения клиенту (5.1–5.2)", "🟢 Высокий эффект", "[  ] Текст в сделку"],
            ["Модуль 6", "Telegram-алерты руководству (6.1–6.3)", "🟢 Критично для РОПа", "[  ] Включить полностью"],
            ["Модуль 7", "Умное движение воронки (7.1–7.2)", "🟡 Опционально", "[  ] Причины отказа"],
            ["Модуль 8", "Персональный AI-коуч менеджера (8.1–8.2)", "🟢 Рекомендуется", "[  ] Личные подсказки в TG"]
        ]
    )
    
    add_callout(doc, 
        "После выбора опций все указанные алгоритмы настраиваются нашими инженерами на вашем сервере "
        "в течение 1 рабочего дня без остановки работы менеджеров в amoCRM. Никаких дополнительных приложений "
        "в компьютер продавцов устанавливать не требуется.",
        title="СРОКИ РАЗВЕРТЫВАНИЯ", alert_type="success")
    
    # Save Word
    output_docx = r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\docs\Каталог_Кастомных_Опций_Интеграции_amoCRM.docx"
    doc.save(output_docx)
    print(f"Word saved to: {output_docx}")
    
    # Copy to presentation
    pres_docx = r"C:\Users\strel\.gemini\antigravity\scratch\revops-enterprise-os\presentation\Каталог_Кастомных_Опций_Интеграции_amoCRM.docx"
    try:
        import shutil
        shutil.copyfile(output_docx, pres_docx)
        print(f"Copied to: {pres_docx}")
    except Exception as e:
        print(f"Notice: {e}")
        
    return output_docx

if __name__ == "__main__":
    generate_custom_features_catalog()
