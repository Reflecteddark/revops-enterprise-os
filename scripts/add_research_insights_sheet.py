"""
Скрипт добавления листа '📑 Инсайты_2_Исследований' в рабочую книгу Фаундера.
Синхронизирует данные с локальным файлом (и на Рабочем столе) и с Google Таблицей.
"""

from __future__ import annotations

import datetime
import os
import shutil
import sys
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import gspread

LOCAL_PLANNER = Path("presentation/RevOps_Founder_Workspace_Planner.xlsx")
DOCS_PLANNER = Path("docs/RevOps_Founder_Workspace_Planner.xlsx")
DESKTOP_PLANNER = Path(os.path.expanduser(r"~\Desktop\RevOps_Founder_Workspace_Planner.xlsx"))

SPREADSHEET_ID = "1AbaQGHhzar7Qsb9U9TByN57rAmOn36zw4HOD2XVrmTc"
CREDENTIALS_FILE = Path("service_account.json")

RESEARCH_DATA = [
    # ── Блок 1: Исследование 1 (43 стр.) ──
    (
        "1.1",
        "Кризис B2B-продаж и рост CAC",
        "Стоимость привлечения лида (CAC) в РФ выросла в 2.5–4 раза. Яндекс Директ перегрет, холодные звонки дают <1% конверсии.",
        "Бизнес больше не может расти экстенсивно за счет вливания бюджета в рекламу. Слив каждого лида обходится в 2 000 – 15 000 ₽.",
        "Фокус RevOps OS: не заливать новый трафик, а ликвидировать скрытые дыры в воронке, возвращая до 35% зависших сделок.",
        "ГЛАВНЫЙ ПИТЧ: «Не тратьте еще 500к на рекламу — мы вернем вам от 1.5М ₽ из текущей базы лидов за 1-й месяц»."
    ),
    (
        "1.2",
        "Уровень утечек 1: «Справочное бюро»",
        "35% лидов сливаются на 1-м звонке: менеджеры подробно отвечают на вопросы клиента, но не берут инициативу и не закрывают на дату.",
        "Клиент получает бесплатную консультацию и уходит фразой «я подумаю». Менеджер ставит статус 'думает' и забывает.",
        "ИИ-Супервайзер по 13 критериям: жесткая проверка CR1 (Next Step). Если точная дата/время не названы — алерт РОПу через 15 минут.",
        "ПРАВИЛО №1: Любой контакт без согласованной даты следующего звонка считается СЛИВOM сделки."
    ),
    (
        "1.3",
        "Уровень утечек 2: «Слепота РОПа»",
        "РОП слушает менее 2% звонков физически (при объеме 300–500 звонков в неделю). Выводы строятся на субъективных ощущениях.",
        "Тренинги по продажам не работают, потому что РОП не видит системных ошибок и не контролирует закрепление навыков.",
        "Telegram-Шериф: 3 персонализированные сводки в день (08:30 утренний фокус, 14:00 SOS-алерты, 18:30 итог дня). Время РОПа: 15 мин/день.",
        "ПРАВИЛО №2: РОПу не нужны графики на 100 страниц — ему нужен список из 3 конкретных сделок, которые горят прямо сейчас."
    ),
    (
        "1.4",
        "Уровень утечек 3: «Кладбище зомби-сделок»",
        "В CRM висят миллионы на этапах «КП отправлено» или «В работе». При реальном аудите 65–75% из них брошены без касания 14+ дней.",
        "Отдел продаж создает иллюзию огромного пайплайна, а фактическая касса пуста. Собственник не видит реального прогноза выручки.",
        "Автоматический Радар Выручки: маркировка зомби-сделок, расчет упущенной выгоды и выгрузка списка реанимации для дожима.",
        "ПРАВИЛО №3: Реанимация 35% зависших сделок в первый месяц окупает подписку RevOps за 3 дня."
    ),
    (
        "1.5",
        "Позиционирование: RevOps-консультант",
        "SaaS-платформы (MANGO, UIS) дают сырые транскрипты без действий. Классический консалтинг стоит от 300к ₽ и длится 2-3 месяца.",
        "Клиенту не нужен голый софт и не нужны длинные отчеты — ему нужны деньги в кассе и внедренные регламенты без затягивания.",
        "Гибрид 'Софт + Консалтинг': 7-дневный пилотный спринт за 29 000 ₽ + регулярная супервизия с гарантией окупаемости.",
        "ПРАВИЛО №4: «Мы не продаем софт. Мы находим, где вы теряете деньги, и пишем пошаговый план, как их вернуть»."
    ),

    # ── Блок 2: Исследование 2 (15 стр.) ──
    (
        "2.1",
        "Революция Unit-экономики STT",
        "Yandex SpeechKit Deferred (Batch STT) стоит ~0.02 ₽/мин ($0.0003), что в 3x дешевле Groq Whisper (~0.06 ₽/мин) и на 100% в РФ.",
        "Себестоимость распознавания всех звонков компании составляет всего 300–600 ₽ в месяц! Маржинальность нашего сервиса > 90%.",
        "Двухуровневая архитектура: Yandex Deferred для ночной 100% обработки звонков + Groq Whisper для быстрого (<5 сек) экспресс-аудита.",
        "ПРАВИЛО №5: Покупка своего сервера с RTX 4090 (~250к ₽) окупается только от 410 000 мин/мес. Облачный Yandex Deferred вне конкуренции."
    ),
    (
        "2.2",
        "Лимиты API amoCRM и Битрикс24",
        "amoCRM режет запросы при > 7 rps. Облачный Битрикс24 жестко блокирует при > 2 rps (пауза 0.5s). Вебхуки Битрикс24 лагают на 5-20 мин.",
        "Без встроенного рейт-лимитера и очереди выгрузка 500 звонков приводит к ошибкам 429/503 и полному падению синхронизации.",
        "Внедрен RateLimiter (6 rps amo, 2 rps Б24), экспоненциальный бэкофф с джиттером и Batch API (до 50 операций в 1 запросе к Б24).",
        "ПРАВИЛО №6: Строго выдерживать паузу не менее 0.51s между вызовами REST API Битрикс24 и группировать запросы в batch."
    ),
    (
        "2.3",
        "Защита от псевдодублей в CRM",
        "Клиенты звонят с разных номеров. Метод crm.duplicate.findbycomm в Битрикс24 возвращает максимум 20 записей. Менеджеры плодят дубли.",
        "Если прикрепить звонок к неверной закрытой сделке или старому дублю, РОП видит искаженную картину и теряет нить переговоров.",
        "Алгоритм resolve_target_deal: нормализация номеров по E.164, приоритет открытых стадий воронки > наибольшая сумма > свежая дата.",
        "ПРАВИЛО №7: Никогда не брать первую попавшуюся сделку из API. Всегда фильтровать по активности и сумме."
    ),
    (
        "2.4",
        "Партнерство с интеграторами CRM",
        "Интеграторы amoCRM и Битрикс24 зарабатывают на разовой настройке и страдают от оттока. Им жизненно необходим recurring-доход.",
        "Интеграторы уже имеют доверие сотен лояльных B2B-клиентов. Партнерская сеть — главный бесплатный канал лидогенерации.",
        "Оффер интеграторам: 25–30% пожизненного ревшара (от 12 250 до 36 000 ₽/мес на каждого клиента). Мы ведем 100% саппорта.",
        "ПРАВИЛО №8: Интегратор не должен ничего внедрять сам. Он просто дает контакт клиента, а мы ежемесячно выплачиваем ему ревшар."
    ),
    (
        "2.5",
        "Контур безопасности 152-ФЗ РФ",
        "Штрафы Роскомнадзора за нарушение 152-ФЗ и хранение ПДн за рубежом достигают 6–18 млн ₽. Любая СБ блокирует сервис без документов.",
        "Клиенты боятся проверок, а менеджеры могут пожаловаться в трудовую инспекцию на прослушку звонков.",
        "Модуль PIISanitizer (токенизация телефонов/email) + комплект из 5 документов (согласие сотрудника, приказ, регламент 30 дней, чек НПД).",
        "ПРАВИЛО №9: Юридический комплект снимает 100% страхов СБ и юристов клиента. Передавать его бесплатно при подключении пилота."
    ),
]


def update_local_excel():
    """Добавляет лист в локальный файл Excel."""
    wb = openpyxl.load_workbook(LOCAL_PLANNER)

    sheet_name = "📑 Инсайты_2_Исследований"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    # Создаем лист
    ws = wb.create_sheet(sheet_name, index=2)
    ws.views.sheetView[0].showGridLines = True

    # Цвета и стили
    c_navy_dark = "0F172A"
    c_navy_light = "1E293B"
    c_border = "CBD5E1"

    fill_navy = PatternFill(start_color=c_navy_dark, end_color=c_navy_dark, fill_type="solid")
    fill_header = PatternFill(start_color=c_navy_light, end_color=c_navy_light, fill_type="solid")
    fill_b1 = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    fill_b2 = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    fill_rule = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")

    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    font_reg = Font(name="Segoe UI", size=10, color="334155")
    font_rule = Font(name="Segoe UI", size=10, bold=True, color="B45309")

    thin_border = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border)
    )

    # Заголовок
    ws.merge_cells("A1:F1")
    ws["A1"] = "Главные Тезисы и Инсайты 2-х Исследований Рынка РФ (43 стр. + 15 стр.) • Что нельзя упустить"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = [
        "№", "Тема / Направление", "Ключевой инсайт исследования (Тезис)",
        "Почему это критически важно (Боль)", "Решение в RevOps OS (Что делаем)",
        "Золотое правило фаундера (Что нельзя упустить)"
    ]

    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(3, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 28

    curr_row = 4
    for item in RESEARCH_DATA:
        num, title, insight, why, sol, rule = item
        ws.row_dimensions[curr_row].height = 54

        ws.cell(curr_row, 1, num).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(curr_row, 1).font = font_bold

        ws.cell(curr_row, 2, title).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(curr_row, 2).font = font_bold

        ws.cell(curr_row, 3, insight).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(curr_row, 3).font = font_reg

        ws.cell(curr_row, 4, why).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(curr_row, 4).font = font_reg

        ws.cell(curr_row, 5, sol).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(curr_row, 5).font = font_bold
        ws.cell(curr_row, 5).fill = fill_b2 if num.startswith("2") else fill_b1

        ws.cell(curr_row, 6, rule).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(curr_row, 6).font = font_rule
        ws.cell(curr_row, 6).fill = fill_rule

        for c in range(1, 7):
            ws.cell(curr_row, c).border = thin_border

        curr_row += 1

    # Настройка ширины колонок
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 38
    ws.column_dimensions["F"].width = 44

    wb.save(LOCAL_PLANNER)
    shutil.copyfile(LOCAL_PLANNER, DOCS_PLANNER)
    if DESKTOP_PLANNER.parent.exists():
        shutil.copyfile(LOCAL_PLANNER, DESKTOP_PLANNER)

    print("Локальный файл Excel успешно обновлен и скопирован на Рабочий стол!")


def update_google_sheets():
    """Синхронизирует новый лист напрямую в Google Таблицу."""
    if not CREDENTIALS_FILE.exists():
        print("service_account.json не найден, пропускаем Google Sheets.")
        return

    gc = gspread.service_account(filename=str(CREDENTIALS_FILE))
    sh = gc.open_by_key(SPREADSHEET_ID)

    sheet_title = "📑 Инсайты_2_Исследований"
    try:
        ws = sh.worksheet(sheet_title)
        sh.del_worksheet(ws)
    except Exception:
        pass

    # Создаем лист
    ws = sh.add_worksheet(title=sheet_title, rows=30, cols=8, index=2)

    # Заполняем данными
    rows_to_insert = [
        ["Главные Тезисы и Инсайты 2-х Исследований Рынка РФ (43 стр. + 15 стр.) • Что нельзя упустить", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        [
            "№", "Тема / Направление", "Ключевой инсайт исследования (Тезис)",
            "Почему это критически важно (Боль)", "Решение в RevOps OS (Что делаем)",
            "Золотое правило фаундера (Что нельзя упустить)"
        ]
    ]

    for item in RESEARCH_DATA:
        rows_to_insert.append(list(item))

    ws.update(range_name=f"A1:F{len(rows_to_insert)}", values=rows_to_insert)
    print("Google Таблица успешно синхронизирована с новым листом!")


if __name__ == "__main__":
    update_local_excel()
    update_google_sheets()
