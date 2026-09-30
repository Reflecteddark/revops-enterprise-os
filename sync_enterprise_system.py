import json
import sys
sys.stdout.reconfigure(encoding='utf-8')
import gspread
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from google.oauth2.service_account import Credentials

print('=== SYNCING ENTERPRISE MULTI-USER RBAC TO GS & XLSX ===')

with open('revops_rbac_v176_hardened.js', 'r', encoding='utf-8') as f:
    code_content = f.read()

# 1. Update Google Sheets Code_Archive
with open('service_account.json', 'r', encoding='utf-8') as f:
    sa = json.load(f)

creds = Credentials.from_service_account_info({
    'type': 'service_account',
    'client_email': sa['email'],
    'private_key': sa['privateKey'],
    'token_uri': 'https://oauth2.googleapis.com/token'
}, scopes=['https://www.googleapis.com/auth/spreadsheets'])

gc = gspread.authorize(creds)
SPREADSHEET_ID = '1QnjrrbpqhYssofchee7G06szWrFvBCcOqGIVokjqdVc'
sh = gc.open_by_key(SPREADSHEET_ID)

ws_code = sh.worksheet('🔐 Code_Archive')
ws_code.clear()
lines = [[l] for l in code_content.split('\n')]
ws_code.update(lines, 'A1:A' + str(len(lines)), value_input_option='RAW')
print(f'Uploaded {len(lines)} lines to 🔐 Code_Archive in Google Sheets.')

# 2. Update XLSX Workbooks
xlsx_files = [
    r'C:\Users\strel\.gemini\antigravity\snapshots\RevOps Platform V17.6 (RBAC Production Suite).xlsx',
    r'C:\Users\strel\.gemini\antigravity\scratch\RevOps Platform V17.6 (RBAC Production Suite).xlsx',
    r'C:\Users\strel\.gemini\antigravity\snapshots\RevOps Platform V17.5 (Operational Wave).xlsx',
    r'C:\Users\strel\.gemini\antigravity\scratch\RevOps Platform V17.5 (Operational Wave).xlsx',
    r'C:\Users\strel\.gemini\antigravity\snapshots\RevOps_Enterprise_OS_V17.6_RBAC_Matrix_Snapshot_2026-W40_20260930_003000.xlsx'
]

thin = Side(border_style='thin', color='CBD5E1')
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

for p in xlsx_files:
    wb = openpyxl.load_workbook(p, data_only=False)
    
    # Update Code_Archive
    if '🔐 Code_Archive' in wb.sheetnames:
        del wb['🔐 Code_Archive']
    ws = wb.create_sheet('🔐 Code_Archive')
    for idx, l in enumerate(code_content.split('\n'), start=1):
        ws.cell(idx, 1, l)
    
    # Update Settings G9:H10
    st = wb['⚙️ Настройки']
    st['G9'] = 'СЕССИЯ ПЛАНИРОВАНИЯ:'
    st['H9'] = '🟢 Готова к работе (Свободна)'
    st['G10'] = 'РЕЖИМ ЭКСПЛУАТАЦИИ:'
    st['H10'] = '👥 Multi-User (Персональные фильтры)'
    
    # Style G9:H10
    for cell_id, bg, fg, align in [
        ('G9', '1E283D', 'FFFFFF', 'right'),
        ('G10', '1E283D', 'FFFFFF', 'right'),
        ('H9', 'E6F4EA', '137333', 'center'),
        ('H10', 'E8F0FE', '1A73E8', 'center')
    ]:
        c = st[cell_id]
        c.font = Font(name='Segoe UI', size=9, bold=True, color=fg)
        c.fill = PatternFill(start_color=bg, end_color=bg, fill_type='solid')
        c.alignment = Alignment(horizontal=align, vertical='center')
        c.border = border_all
    
    dv_mode = DataValidation(type='list', formula1='"👥 Multi-User (Персональные фильтры),🎯 Solo Focus (Автоскрытие листов)"', allow_blank=False)
    st.add_data_validation(dv_mode)
    dv_mode.add('H10')

    wb.save(p)
    print(f'Synced XLSX: {p}')

# 3. Verify QA
ws_qa = sh.worksheet('🧪 QA_Suite')
res = ws_qa.batch_get(['B4', 'C4', 'D4', 'A3'])
print('QA Status in Google Sheets:', res)
